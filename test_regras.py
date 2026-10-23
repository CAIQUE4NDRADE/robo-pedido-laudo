"""Testes das regras de negócio do robô (sem navegador, sem Outlook, sem SharePoint)."""
import io
import queue
from datetime import datetime, timedelta

import openpyxl
import pytest

import monitor_laudo as m


# ---------------------------------------------------------------- leitura do email de virada
@pytest.mark.parametrize("assunto, esperado", [
    ("RE: [Externo] Virada de OD S10 p/ TQ 09 às 04h30 - TODOS", ("OD S10", "09", "04h30")),
    ("[Externo] Virada de GASOLINA p/ TQ-01 às 06h00 - TODOS", ("GASOLINA", "01", "06h00")),
    ("[Externo] Virada de S500 p/ TQ-02 às 08h20 - TODOS", ("S500", "02", "08h20")),
    ("Virada de Gasolina A para o tanque 103 as 14:00", ("GASOLINA A", "103", "14h00")),
])
def test_extrair_virada(assunto, esperado):
    v = m.extrair_virada(assunto)
    assert (v["produto"], v["tanque"], v["hora"]) == esperado


@pytest.mark.parametrize("texto", [
    "Reunião amanhã",
    "Solicitação de laudo - TQ 09 (Pool) + TQ 103 (Base) - Houve virada de OD S10 para o TQ 09",
    "Microsoft Outlook RE: [Externo] Virada de OD S10 p/ TQ 09 às 04h30 - Não foi possível entregar",
])
def test_nao_e_virada(texto):
    assert m.extrair_virada(texto) is None


def test_composto():
    assert m.produto_base_para("OD S10") == "BIODIESEL B100"
    assert m.produto_base_para("S500") == "BIODIESEL B100"
    assert m.produto_base_para("GAS") == "ETANOL ANIDRO"
    assert m.produto_base_para("JET A") is None


# ---------------------------------------------------------------- datas e turnos
@pytest.mark.parametrize("hora, agora, esperado", [
    ("21h00", datetime(2026, 10, 19, 1, 45), datetime(2026, 10, 18, 21, 0)),   # ontem
    ("06h30", datetime(2026, 10, 19, 6, 29), datetime(2026, 10, 19, 6, 30)),   # aviso antes da virada
    ("00h10", datetime(2026, 10, 19, 23, 50), datetime(2026, 10, 20, 0, 10)),  # amanhã
])
def test_data_virada(hora, agora, esperado):
    assert m.data_virada({"hora": hora}, agora) == esperado


@pytest.mark.parametrize("hm, turno", [((3, 44), "3"), ((7, 29), "3"), ((7, 30), "1"), ((15, 30), "2")])
def test_turno(hm, turno):
    assert m.turno_de(datetime(2026, 10, 19, *hm)) == turno


def test_item_de_hoje():
    assert m.item_de_hoje("Virada de OD S500 p/ TQ 04 Dom, 21:08") is False
    assert m.item_de_hoje("Virada de GAS p/ TQ 06 Ontem 15:16") is False
    agora = datetime.now()
    assert m.item_de_hoje(f"Virada de OD S10 p/ TQ 09 {agora:%H:%M}") is True


# ---------------------------------------------------------------- planilha de passagem de turno
def _passagem(t3, t1=None, t2=None, laudos=(), aba=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = aba or str(datetime.now().day)
    ws["A4"] = "PRODUTO"
    ws["B4"], ws["G4"], ws["L4"] = "3° TURNO", "1° TURNO", "2° TURNO"
    for c in "BGL":
        ws[f"{c}5"] = "EXPEDINDO"
    for i, prod in enumerate(["GAS A", "B100", "EA", "EH"]):
        r = 6 + i
        ws[f"A{r}"] = prod
        for col, vals in (("B", t3), ("G", t1 or {}), ("L", t2 or {})):
            ws[f"{col}{r}"] = vals.get(prod)
    ws["A15"], ws["G15"], ws["I15"] = "Tanque", "Amostra Retirada", "Laudo Emitido"
    for i, (tq, am, em) in enumerate(laudos):
        ws[f"A{16 + i}"], ws[f"G{16 + i}"], ws[f"I{16 + i}"] = tq, am, em
    ws["A26"] = "ADIAMENTO DE MANUTENÇÃO"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_passagem_usa_turno_atual_e_ignora_turno_futuro():
    dados = _passagem({"B100": "203", "EA": "202"}, t1={"B100": "105", "EA": "102"}, aba="19")
    p = m.ler_passagem(dados, datetime(2026, 10, 19, 3, 44))     # 3º turno: 1º ainda não começou
    assert p["BIODIESEL B100"]["valor"] == "203"
    p = m.ler_passagem(dados, datetime(2026, 10, 19, 8, 10))     # 1º turno preenchido
    assert p["BIODIESEL B100"]["valor"] == "105"


@pytest.mark.parametrize("valor, esperado", [
    ("203", ["203"]), ("206/105", ["105"]), ("206", ["206"]), ("203/206", ["203", "206"]), ("103", ["103"]),
])
def test_regra_b100(valor, esperado):
    p = {"BIODIESEL B100": {"valor": valor, "turno": "1", "aba": "19"}}
    assert m.tanques_passagem("BIODIESEL B100", p) == esperado


def test_laudos_hidratado():
    dados = _passagem({}, laudos=[("117", "Sim", ""), ("105", "Sim", ""), ("212", "", "")])
    linhas = m.ler_laudos_passagem(dados)
    assert [(ln["tanque"], ln["amostra"]) for ln in linhas] == [("117", "Sim"), ("212", "")]


def test_abertura():
    texto = "19/10/2026\tGAS\t6\t(2)INAM\n19/10/2026\tS10\t5\tAMARELADO\n19/10/2026\tS500\t4\tVERMELHO\n19/10/2026\tOC1A\t11"
    assert m.parse_abertura(texto) == {"GAS": "6", "S10": "5", "S500": "4"}


# ---------------------------------------------------------------- monitor (envio simulado)
@pytest.fixture
def monitor(monkeypatch):
    enviados = []
    monkeypatch.setattr(m, "enviar_email", lambda d, para, cc, assunto, corpo, log=None: enviados.append((para, assunto, corpo)))
    monkeypatch.setattr(m, "salvar_estado", lambda e: None)
    mon = m.Monitor(queue.Queue(), modo_teste=True)
    mon.estado = {"chaves": {}, "historico": [], "enviadas": {}}
    mon.enviados = enviados
    return mon


def test_troca_b100_vai_com_ultimo_pool_e_nao_duplica(monitor):
    agora = datetime.now()
    monitor.estado["aberturas"] = {agora.strftime("%d/%m/%Y"): {"S10": "5", "S500": "4"}}
    monitor.registrar_pool("OD S10", "5", agora - timedelta(hours=3))
    monitor.registrar_pool("S500", "4", agora - timedelta(hours=3))
    P = lambda v: {"BIODIESEL B100": {"valor": v, "turno": "1", "aba": "19"}}
    monitor.passagem = P("103")
    monitor.verificar_trocas_base(inicial=True)
    assert monitor.enviados == []                       # ao ligar só memoriza
    monitor.passagem = P("105")
    monitor.verificar_trocas_base()
    monitor.verificar_trocas_base()
    assert len(monitor.enviados) == 1                   # troca 103 -> 105, uma vez
    corpo = monitor.enviados[0][2]
    assert "TQ 105 (Base, em expedição)" in corpo and "TQ 05 (Pool) - S10" in corpo and "TQ 04 (Pool) - S500" in corpo


def test_virada_sem_tanque_da_base_nao_envia(monitor):
    monitor.passagem = {}
    monitor.atualizar_passagem = lambda: None
    v = m.extrair_virada("Virada de GAS p/ TQ 01 às 06h30")
    monitor.processar(v, "conv1|GAS|1|06h30")
    assert monitor.enviados == [] and "conv1|GAS|1|06h30" in monitor.pendentes


def test_semanal_s10_b30_so_na_janela(monitor):
    monitor.atualizar_passagem = lambda: setattr(monitor, "passagem", {"BIODIESEL B100": {"valor": "203", "turno": "2", "aba": "25"}})
    monitor.garantir_abertura_hoje = lambda: None
    monitor.registrar_pool("OD S10", "5", datetime(2026, 10, 25, 22, 0))
    for quando in [datetime(2026, 10, 26, 0, 20), datetime(2026, 10, 26, 0, 31),
                   datetime(2026, 10, 26, 0, 45), datetime(2026, 10, 26, 1, 5)]:
        monitor.verificar_semanal_s10b30(quando)
    assert len(monitor.enviados) == 1 and "S10 B30" in monitor.enviados[0][1]
