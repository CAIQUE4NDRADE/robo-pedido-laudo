# -*- coding: utf-8 -*-
"""
Robô de Pedido de Laudo - Terminal de combustíveis
=================================================

Automatiza a solicitação de laudos de qualidade ao laboratório. Só pede laudo quando
existe um fato operacional que exige (antes os laudos eram gerados para todos os tanques
dos compostos, sem solicitação, com custo alto).

Gatilhos:
  1. Email de VIRADA de tanque do pool (Outlook Web, busca "virada")
       -> laudo do tanque do pool + tanque da base do componente do composto
          (S10 / S500 -> Biodiesel B100   |   Gasolina -> Etanol Anidro)
  2. TROCA do tanque de B100 / Anidro em expedição (planilha de Passagem de Turno, SharePoint)
       -> laudo do tanque novo + último tanque do pool de cada produto que usa o componente
  3. Email de ABERTURA DE TANQUES (noite) com tanque de S10/S500/Gasolina diferente do dia anterior
  4. Etanol Hidratado lançado na seção LAUDOS da Passagem de Turno (amostra retirada)
  5. Pedido semanal (segunda 00:30-01:00): S10 B30
  + Aviso por email ao responsável se a abertura de tanques não chegar até 00:15.

Proteções: só processa emails que chegam depois de o robô ser ligado, nunca envia duplicado,
nunca envia pedido incompleto (segura e avisa), modo teste, rótulo de confidencialidade automático.

Configuração da instalação: copie config.example.json para config.json (não versionado).
Requisitos:  pip install -r requirements.txt
"""
import json
import queue
import re
import socket
import threading
import time
import traceback
import urllib.parse
from datetime import datetime, timedelta
from pathlib import Path

import tkinter as tk
from tkinter import ttk, messagebox

from selenium import webdriver
from selenium.common.exceptions import StaleElementReferenceException
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By

# ============================== CONFIGURAÇÃO ==============================
BASE = Path(__file__).parent

# config.json (não versionado) sobrescreve os valores de exemplo abaixo
try:
    _CONFIG = json.loads((BASE / "config.json").read_text(encoding="utf-8"))
except Exception:
    _CONFIG = {}


def _cfg(nome, padrao):
    v = _CONFIG.get(nome, padrao)
    if isinstance(padrao, tuple) and isinstance(v, list):
        v = tuple(v)
    return v


PERFIL_CHROME = BASE / "chrome_perfil_virada"
ARQ_ESTADO = BASE / "viradas_processadas.json"
ARQ_LOG = BASE / "log_virada.txt"

EMAIL_LABORATORIO = _cfg("EMAIL_LABORATORIO", "laboratorio@exemplo.com")
EMAIL_CC = _cfg("EMAIL_CC", [])               # ex.: ["supervisao@exemplo.com"]

# MODO TESTE: manda para o responsável no lugar do laboratório (sem CC).
# Para produção: MODO_TESTE = False (ou desmarque a caixa na tela).
MODO_TESTE = _cfg("MODO_TESTE", True)
EMAIL_TESTE = _cfg("EMAIL_TESTE", "responsavel@exemplo.com")
ASSINATURA = ""                   # vazio = usa a assinatura automática do Outlook

OUTLOOK_INBOX = "https://outlook.office.com/mail/inbox"
OUTLOOK_COMPOSE_PATH = "/mail/deeplink/compose"   # usa o domínio que o Outlook abrir (cloud.microsoft)
ROTULO_CONFIDENCIALIDADE = "Público"   # escolhido quando o Outlook pede rótulo ao enviar
TERMO_BUSCA = "virada"                # busca do Outlook (Todas as pastas)
HORAS_EMAIL_X_VIRADA = 3              # email só vale se chegou até 3h do horário da virada
RETENTAR_PENDENTE_S = 120             # virada sem tanque Base: tenta de novo a cada 2 min
MAX_TENTATIVAS = 3                 # por virada; depois desiste e avisa no log
HORAS_MESMA_VIRADA = 20            # RE:/encaminhado da mesma virada nesse prazo não reenvia
RE_NAO_VIRADA = re.compile(r"n[ãa]o foi poss[íi]vel entregar|undeliverable|falha na entrega|delivery has failed|mail delivery", re.IGNORECASE)
# o próprio pedido de laudo (ou resposta a ele) cita "virada de ... para o TQ" no texto: nunca é virada
RE_PEDIDO_LAUDO = re.compile(r"solicita[çc][ãa]o de laudo", re.IGNORECASE)

INTERVALO_S = 20                   # leitura da caixa de entrada
RECARREGAR_A_CADA_S = 600          # F5 de segurança na caixa
DIAS_GUARDAR = 30                  # quanto tempo lembrar viradas já enviadas

# Tanques Base de cada componente (exemplo - ajuste no config.json)
TANQUES_BASE = _cfg("TANQUES_BASE", {
    "BIODIESEL B100": ["103", "105", "203", "206"],
    "ETANOL ANIDRO":  ["101", "102", "116", "201", "202"],
})
ARQ_EXPEDINDO = BASE / "tanques_expedindo.json"   # reserva escolhida na tela

# Passagem de Turno (SharePoint) - fonte principal do tanque expedindo
SP_SITE = _cfg("SP_SITE", "https://SUAEMPRESA.sharepoint.com/sites/SEU_SITE")
# Pasta da passagem de turno no SharePoint
# Estrutura: <SP_RAIZ> / <ano> / "10 - Outubro" / Passagem de Turno ... A1 e A2.xlsx
SP_RAIZ = _cfg("SP_RAIZ", "/sites/SEU_SITE/Documentos Compartilhados/OPERACOES/PASSAGEM DE TURNO")
SP_PASTA = f"{SP_RAIZ}/{datetime.now().year}"
SP_PASTA_URL = (f"{SP_SITE}/Documentos%20Compartilhados/Forms/AllItems.aspx?id="
                + urllib.parse.quote(SP_PASTA, safe=""))
NOME_PLANILHA = _cfg("NOME_PLANILHA", "A1 E A2")   # parte do nome da planilha (a de DESVIO é ignorada)
# arquivo usado se não der para listar a pasta
SP_ARQUIVO_ID = _cfg("SP_ARQUIVO_ID", "00000000-0000-0000-0000-000000000000")  # planilha padrão (reserva)
SP_DOWNLOAD = f"{SP_SITE}/_layouts/15/download.aspx?UniqueId={SP_ARQUIVO_ID}"
MESES = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
INTERVALO_PASSAGEM_S = 120          # relê a planilha a cada 2 min (e sempre na hora da virada)
B100_PREFERIDOS = _cfg("B100_PREFERIDOS", ("103", "105"))    # B100: se a passagem citar 103/105, usa só eles; senão usa o que estiver (203/206)
LINHA_PASSAGEM = {"BIODIESEL B100": ("B100",), "ETANOL ANIDRO": ("EA", "ETANOL ANIDRO", "ANIDRO"),
                  "ETANOL HIDRATADO": ("EH", "ETANOL HIDRATADO", "HIDRATADO")}
# troca do tanque Base em expedição (na passagem) -> pede laudo do tanque novo
PRODUTOS_TROCA = ("BIODIESEL B100", "ETANOL ANIDRO")
# Hidratado: pede laudo quando o supervisor lança o tanque na seção LAUDOS da passagem
# (coluna "Tanque" + "Amostra Retirada" preenchidas e "Laudo Emitido" ainda vazio)
TANQUES_HIDRATADO = _cfg("TANQUES_HIDRATADO", ("117", "121", "212", "213"))
# Troca de B100 / Anidro: o pedido vai com o ÚLTIMO tanque do pool de cada produto que usa o componente
# Pedido semanal: toda SEGUNDA às 00:30 -> laudo do S10 B30 (último TQ de S10 do pool + B100 expedindo).
# Janela de envio: entre 00:30 e 01:00 (fora disso não envia naquela semana).
SEMANAL_DIA = 0                       # 0 = segunda-feira
SEMANAL_HORA = (0, 30)
SEMANAL_ATE = (1, 0)
POOL_DO_COMPONENTE = {"BIODIESEL B100": ("S10", "S500"), "ETANOL ANIDRO": ("GAS",)}

# Email "[Externo] ABERTURA DE TANQUES - dd/mm/aaaa" (pool): lido entre 23:25 e 01:00.
# Se o tanque de S10 / S500 / Gasolina mudou em relação à abertura do dia anterior -> pede laudo.
TERMO_ABERTURA = "ABERTURA DE TANQUES"
JANELA_ABERTURA = ((23, 25), (1, 0))
# Se às 00:15 o robô ainda não tiver a abertura do dia, avisa por email:
EMAIL_AVISO_ABERTURA = _cfg("EMAIL_AVISO_ABERTURA", "responsavel@exemplo.com")
AVISO_ABERTURA_HORA = (0, 15)
INTERVALO_ABERTURA_S = 180
INTERVALO_ABERTURA_FORA_S = 600       # fora da janela: se a abertura de hoje ainda não foi lida (ex.: chegou
                                      # encaminhada de manhã), procura de novo a cada 10 min
PRODUTOS_ABERTURA = {"GAS": "GAS", "GASOLINA": "GAS", "S10": "OD S10", "S500": "OD S500"}
RE_LINHA_ABERTURA = re.compile(r"(\d{2}/\d{2}/\d{4})\s+(GASOLINA|GAS|S10|S500)\s+(\d{1,3})\b", re.IGNORECASE)
ORDEM_TURNOS = ("2", "1", "3")      # do mais recente para o mais antigo dentro da aba do dia (3º->1º->2º)
# Horário dos turnos (todos dentro do mesmo dia / mesma aba):
#   3º 00:00-07:30   |   1º 07:30-15:30   |   2º 15:30-24:00
INICIO_1_TURNO = (7, 30)
INICIO_2_TURNO = (15, 30)


def turno_de(quando):
    hm = (quando.hour, quando.minute)
    if hm < INICIO_1_TURNO:
        return "3"
    if hm < INICIO_2_TURNO:
        return "1"
    return "2"


def turnos_ate(quando):
    """Turno do horário e os anteriores do mesmo dia, do mais recente para o mais antigo.
    Nunca olha turno que ainda não começou."""
    return {"3": ("3",), "1": ("1", "3"), "2": ("2", "1", "3")}[turno_de(quando)]

# Produto da virada (regex) -> produto do tanque Base cujo laudo vai junto
REGRAS_COMPOSTO = [
    (r"\bOD\b|DIESEL|S10|S500", "BIODIESEL B100"),   # OD S10 / OD S500 -> biodiesel
    (r"\bGAS",                   "ETANOL ANIDRO"),    # GAS / GASOLINA   -> etanol anidro
]
# ==========================================================================

RE_VIRADA = re.compile(
    r"virada\s+de\s+(?P<produto>.+?)\s+p(?:ara|/)?\s*(?:o\s+)?t(?:q|anque)[\s.\-:#nº°]*(?P<tanque>\d{1,4})"
    r"(?:\s*(?:às|as|a)\s*(?P<hora>\d{1,2}\s*[h:]\s*\d{0,2}))?",
    re.IGNORECASE,
)


# ------------------------------ regras ------------------------------------
def normaliza_hora(h):
    if not h:
        return ""
    n = re.findall(r"\d+", h)
    hh = int(n[0])
    mm = int(n[1]) if len(n) > 1 and n[1] else 0
    return f"{hh:02d}h{mm:02d}"


def extrair_virada(texto):
    t = (texto or "").lower()
    if RE_PEDIDO_LAUDO.search(t):
        return None
    if RE_NAO_VIRADA.search(t) and "microsoft outlook;" not in t and "; microsoft outlook" not in t:
        return None   # item que é só o aviso de "não foi possível entregar"
    m = RE_VIRADA.search(texto or "")
    if not m:
        return None
    return {
        "produto": " ".join(m.group("produto").split()).upper(),
        "tanque": m.group("tanque"),
        "hora": normaliza_hora(m.group("hora")),
    }


RE_DATA_HORA = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\s+(?:às\s+)?\d{1,2}:\d{2}\b")
RE_SEMANA_HORA = re.compile(r"\b(seg|ter|qua|qui|sex|s[áa]b|dom)\.?,\s*\d{1,2}:\d{2}\b|\bontem\b", re.IGNORECASE)
RE_SO_HORA = re.compile(r"(?<![\d/:])\b\d{1,2}:\d{2}\b")


def item_de_hoje(txt, hoje=None):
    """Na lista do Outlook: email de hoje mostra só a hora ('01:45'); de dias anteriores mostra
    'Dom, 21:08', 'Ontem' ou a data. Se a descrição trouxer 'dd/mm/aaaa hh:mm', usa ela."""
    hoje = (hoje or datetime.now()).date()
    m = RE_DATA_HORA.search(txt or "")
    if m:
        try:
            return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).date() == hoje
        except ValueError:
            pass
    if RE_SEMANA_HORA.search(txt or ""):
        return False
    return bool(RE_SO_HORA.search(txt or ""))


def hora_do_item(txt, agora=None):
    """Horário de chegada do email na lista do Outlook (só emails de HOJE mostram 'hh:mm').
    Retorna datetime ou None (email de dia anterior / sem horário)."""
    agora = agora or datetime.now()
    if not item_de_hoje(txt, agora):
        return None
    horas = []
    for mt in RE_SO_HORA.finditer(txt or ""):
        hh, mm = (int(x) for x in mt.group(0).split(":"))
        if hh > 23 or mm > 59:
            continue
        dt = agora.replace(hour=hh, minute=mm, second=0, microsecond=0)
        if dt <= agora + timedelta(minutes=1):
            horas.append(dt)
    return max(horas) if horas else None


def produto_base_para(produto_virada):
    for padrao, prod in REGRAS_COMPOSTO:
        if re.search(padrao, produto_virada, re.IGNORECASE):
            return prod
    return None


def carregar_expedindo():
    """{'BIODIESEL B100': {'tanque': '103', 'quando': '12/10 06:15'}, ...}"""
    try:
        return json.loads(ARQ_EXPEDINDO.read_text(encoding="utf-8"))
    except Exception:
        return {}


def salvar_expedindo(d):
    ARQ_EXPEDINDO.write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")


def _txt(v):
    return " ".join(str(v).split()).upper() if v is not None else ""


def ler_aba_passagem(ws, turnos_validos=ORDEM_TURNOS, filtro=None):
    """{'BIODIESEL B100': {'valor': '206/105', 'turno': '1'}, 'ETANOL ANIDRO': {...}}"""
    turnos, expedindo, col_prod, linhas = [], [], None, {}
    for row in ws.iter_rows(min_row=1, max_row=40, max_col=30):
        for c in row:
            t = _txt(c.value)
            if not t:
                continue
            m = re.match(r"^(\d)\s*[º°O]?\s*TURNO", t)
            if m:
                turnos.append((c.column, m.group(1)))
            elif t == "EXPEDINDO":
                expedindo.append((c.row, c.column))
            elif t == "PRODUTO" and col_prod is None:
                col_prod = c.column
    if not turnos or not expedindo or col_prod is None:
        return {}
    turnos.sort()
    col_turno = {}
    for _, col in expedindo:
        dono = [n for cini, n in turnos if cini <= col]
        if dono:
            col_turno[dono[-1]] = col
    for r in range(1, 41):
        rot = _txt(ws.cell(r, col_prod).value)
        for prod, nomes in LINHA_PASSAGEM.items():
            if rot in nomes and prod not in linhas:
                linhas[prod] = r
    out = {}
    for prod, r in linhas.items():
        for n in turnos_validos:
            col = col_turno.get(n)
            if not col:
                continue
            val = _txt(ws.cell(r, col).value)
            if re.search(r"\d{3}", val) and (filtro is None or filtro(prod, val)):
                out[prod] = {"valor": val, "turno": n}
                break
    return out


def ler_passagem(conteudo, agora=None):
    """Lê o .xlsx da passagem (bytes) para o horário 'agora':
    turno atual da aba do dia; se estiver vazio, o turno anterior; depois a aba de ontem."""
    import io
    import openpyxl
    agora = agora or datetime.now()
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=False)
    abas = [(agora, turnos_ate(agora)), (agora - timedelta(days=1), ORDEM_TURNOS)]
    resultado = {}
    for dia, turnos in abas:
        nome = str(dia.day)
        if nome not in wb.sheetnames or dia.month != agora.month:
            continue
        for prod, info in ler_aba_passagem(wb[nome], turnos).items():
            if prod not in resultado:
                info["aba"] = nome
                resultado[prod] = info
        if len(resultado) == len(LINHA_PASSAGEM):
            break

    return resultado


def tanques_passagem(prod, passagem):
    """Tanques do produto Base na passagem (B100: 103/105 se citados). -> ['105', ...]"""
    info = (passagem or {}).get(prod)
    if not info:
        return []
    toks = re.findall(r"\d{3}", info["valor"])
    if prod == "BIODIESEL B100" and any(t in B100_PREFERIDOS for t in toks):
        toks = [t for t in toks if t in B100_PREFERIDOS]
    return list(dict.fromkeys(toks))


def montar_email_livre(assunto, intro, linhas, agora=None):
    agora = agora or datetime.now()
    pedido = "o envio dos laudos dos tanques abaixo" if len(linhas) > 1 else "o envio do laudo do tanque abaixo"
    corpo = (f"{saudacao(agora)},\n\n{intro}\n\nSolicito, por gentileza, {pedido}:\n"
             + "\n".join(linhas) + "\n\n" + "Atenciosamente," + (f"\n{ASSINATURA}" if ASSINATURA else ""))
    return assunto, corpo


def ler_laudos_passagem(conteudo, agora=None):
    """Linhas da seção LAUDOS da aba do dia (e de ontem):
    [{'aba','linha','tanque','amostra','emitido'}] só dos tanques de hidratado."""
    import io
    import openpyxl
    agora = agora or datetime.now()
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=False)
    out = []
    for dia in (agora - timedelta(days=1), agora):
        nome = str(dia.day)
        if nome not in wb.sheetnames or dia.month != agora.month:
            continue
        ws = wb[nome]
        cab = None
        for row in ws.iter_rows(min_row=1, max_row=80, max_col=30):
            textos = {c.column: _txt(c.value) for c in row if c.value is not None}
            if "TANQUE" in textos.values() and any("AMOSTRA" in t for t in textos.values()):
                cab = (row[0].row, textos)
                break
        if not cab:
            continue
        r0, textos = cab
        col_tq = next(c for c, t in textos.items() if t == "TANQUE")
        col_am = next(c for c, t in textos.items() if "AMOSTRA" in t)
        col_em = next((c for c, t in textos.items() if "LAUDO EMITIDO" in t or "EMITIDO" in t), None)
        for r in range(r0 + 1, r0 + 30):
            tq_txt = _txt(ws.cell(r, col_tq).value)
            if any(p in tq_txt for p in ("ADIAMENTO", "CONTROLE DE LACRES", "EQUIPAMENTO")):
                break
            tqs = [t for t in re.findall(r"\d{3}", tq_txt) if t in TANQUES_HIDRATADO]
            if not tqs and not ("EH" in tq_txt.split() or "HIDRAT" in tq_txt):
                continue
            amostra = ws.cell(r, col_am).value
            emitido = ws.cell(r, col_em).value if col_em else None
            for t in tqs:
                out.append({"aba": nome, "linha": r, "tanque": t,
                            "amostra": "" if amostra is None else str(amostra).strip(),
                            "emitido": "" if emitido is None else str(emitido).strip()})
    return out


def produto_pool(produto_virada):
    """'OD S10' -> 'S10', 'S500'/'OD S500' -> 'S500', 'GAS'/'GASOLINA' -> 'GAS'"""
    p_ = (produto_virada or "").upper()
    if "S500" in p_:
        return "S500"
    if "S10" in p_:
        return "S10"
    if "GAS" in p_:
        return "GAS"
    return None


def parse_abertura(texto):
    """Tabela do email de abertura -> {'GAS': '6', 'S10': '5', 'S500': '4'}"""
    out = {}
    for m_ in RE_LINHA_ABERTURA.finditer(texto or ""):
        prod = m_.group(2).upper()
        prod = "GAS" if prod.startswith("GAS") else prod
        out.setdefault(prod, m_.group(3))
    return out


def escolher_tanques(produto_virada, passagem):
    """-> (produto do componente, ['105-BIODIESEL B100', ...], origem)"""
    prod = produto_base_para(produto_virada)
    if not prod:
        return prod, [], ""
    info = (passagem or {}).get(prod)
    if info:
        toks = re.findall(r"\d{3}", info["valor"])
        if prod == "BIODIESEL B100" and any(t in B100_PREFERIDOS for t in toks):
            toks = [t for t in toks if t in B100_PREFERIDOS]   # "206/105" -> 105
        toks = list(dict.fromkeys(toks))
        if toks:
            return prod, [f"{t}-{prod}" for t in toks], \
                f"passagem aba {info['aba']} / {info['turno']}º turno ('{info['valor']}')"
    tq = (carregar_expedindo().get(prod) or {}).get("tanque", "")
    if tq:
        return prod, [f"{tq}-{prod}"], "reserva da tela"
    return prod, [], ""


def numero_tanque(nome):
    return nome.split("-", 1)[0].strip()


def saudacao(agora=None):
    h = (agora or datetime.now()).hour
    if 5 <= h < 12:
        return "Bom dia"
    if 12 <= h < 18:
        return "Boa tarde"
    return "Boa noite"


def data_virada(v, agora=None):
    """Data da virada = o dia (ontem, hoje ou amanhã) em que o horário dela fica mais perto de agora.
    Ex.: 21h00 lido à 01h45 -> ontem; 00h10 avisado às 23h50 -> amanhã; 06h30 avisado às 06h29 -> hoje."""
    agora = agora or datetime.now()
    if not v.get("hora"):
        return agora
    hh, mm = int(v["hora"][:2]), int(v["hora"][3:5])
    base = agora.replace(hour=hh % 24, minute=mm % 60, second=0, microsecond=0)
    return min((base + timedelta(days=d) for d in (-1, 0, 1)), key=lambda dt: abs(dt - agora))


def montar_email(v, tanques_rz, agora=None):
    """tanques_rz: ex. ['103-BIODIESEL B100']"""
    agora = agora or datetime.now()
    dv = data_virada(v, agora)
    hora = f" às {v['hora']} do dia {dv:%d/%m/%Y}" if v["hora"] else f" em {dv:%d/%m/%Y}"

    # sem "- " no começo: o Outlook transformaria em lista automática ao digitar
    linhas = [f"TQ {v['tanque']} (Pool) - {v['produto']}"]
    for nome in tanques_rz:
        prod = nome.split("-", 1)[1].strip() if "-" in nome else ""
        linhas.append(f"TQ {numero_tanque(nome)} (Base, em expedição) - {prod}")

    if tanques_rz:
        rz = " + ".join(f"TQ {numero_tanque(n)}" for n in tanques_rz)
        assunto = f"Solicitação de laudo - TQ {v['tanque']} (Pool) + {rz} (Base) - {v['produto']}"
        pedido = "o envio dos laudos dos tanques abaixo"
    else:
        assunto = f"Solicitação de laudo - TQ {v['tanque']} (Pool) - {v['produto']}"
        pedido = "o envio do laudo do tanque abaixo"

    corpo = (
        f"{saudacao(agora)},\n\n"
        f"Houve virada de {v['produto']} para o TQ {v['tanque']} (Pool){hora}.\n\n"
        f"Solicito, por gentileza, {pedido}:\n"
        + "\n".join(linhas) + "\n\n"
        "Atenciosamente," + (f"\n{ASSINATURA}" if ASSINATURA else "")
    )
    return assunto, corpo


def msg_curta(e):
    """Só a primeira linha da mensagem do Selenium (sem o stacktrace gigante)."""
    return str(e).replace("Message:", "").strip().split("\n")[0][:200]


# ------------------------------ estado ------------------------------------
ARQ_ESTADO_BAK = BASE / "viradas_processadas.bak.json"


def carregar_estado():
    for arq in (ARQ_ESTADO, ARQ_ESTADO_BAK):   # se o principal sumir/corromper, usa a cópia
        if arq.exists():
            try:
                e = json.loads(arq.read_text(encoding="utf-8"))
                e.setdefault("chaves", {})
                e.setdefault("historico", [])
                return e
            except Exception:
                continue
    return None


def salvar_estado(estado):
    limite = (datetime.now() - timedelta(days=DIAS_GUARDAR)).isoformat()
    estado["chaves"] = {k: t for k, t in estado["chaves"].items() if t >= limite}
    estado["enviadas"] = {k: t for k, t in estado.get("enviadas", {}).items() if t >= limite}
    estado["laudos_vistos"] = {k: t for k, t in estado.get("laudos_vistos", {}).items() if t >= limite}
    estado["historico"] = estado["historico"][-500:]
    tmp = ARQ_ESTADO.with_suffix(".tmp")
    tmp.write_text(json.dumps(estado, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(ARQ_ESTADO)
    try:
        ARQ_ESTADO_BAK.write_text(ARQ_ESTADO.read_text(encoding="utf-8"), encoding="utf-8")
    except Exception:
        pass


# ------------------------------ navegador ---------------------------------
def _fechar_chrome_do_perfil(perfil):
    """Fecha só os Chrome que estão usando o NOSSO perfil (sobra de execução anterior).
    Não mexe no Chrome normal do usuário nem nos outros monitores."""
    import subprocess
    nome = Path(perfil).name
    ps = ("Get-CimInstance Win32_Process -Filter \"Name='chrome.exe'\" | "
          f"Where-Object {{ $_.CommandLine -like '*{nome}*' }} | "
          "ForEach-Object { Stop-Process -Id $_.ProcessId -Force -ErrorAction SilentlyContinue }")
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], timeout=30,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except Exception:
        pass
    time.sleep(2)


def _limpar_travas(perfil):
    """Arquivos de trava que ficam quando o Chrome fecha à força e impedem abrir de novo."""
    for rel in ("SingletonLock", "SingletonCookie", "SingletonSocket", "lockfile",
                "Default/LOCK", "Default/Last Session", "Default/Last Tabs"):
        try:
            (Path(perfil) / rel).unlink()
        except Exception:
            pass


def _abrir_chrome(perfil, service):
    opts = Options()
    if perfil:
        Path(perfil).mkdir(parents=True, exist_ok=True)
        opts.add_argument(f"--user-data-dir={perfil}")
        opts.add_argument("--profile-directory=Default")
    opts.add_argument("--start-maximized")
    opts.add_argument("--no-first-run")
    opts.add_argument("--no-default-browser-check")
    opts.add_argument("--disable-session-crashed-bubble")
    # PC robô: não deixar o Chrome "dormir" quando a janela fica atrás de outra ou minimizada
    opts.add_argument("--disable-background-timer-throttling")
    opts.add_argument("--disable-backgrounding-occluded-windows")
    opts.add_argument("--disable-renderer-backgrounding")
    opts.add_experimental_option("excludeSwitches", ["enable-automation"])
    d = webdriver.Chrome(service=service, options=opts)
    d.set_page_load_timeout(60)
    return d


def criar_driver(log=print):
    try:
        from webdriver_manager.chrome import ChromeDriverManager
        socket.setdefaulttimeout(30)
        service = Service(ChromeDriverManager().install())
    except Exception:
        service = Service()  # fallback: Selenium Manager / driver em cache
    finally:
        socket.setdefaulttimeout(None)

    tentativas = [
        ("normal", PERFIL_CHROME, False),
        ("fechando Chrome antigo do perfil e limpando travas", PERFIL_CHROME, True),
        ("perfil reserva (vai pedir login de novo)", BASE / "chrome_perfil_virada_2", True),
    ]
    ultimo = None
    for desc, perfil, limpar in tentativas:
        if limpar:
            log(f"Chrome não abriu - tentando de novo ({desc})...")
            _fechar_chrome_do_perfil(perfil)
            _limpar_travas(perfil)
        try:
            return _abrir_chrome(perfil, service)
        except Exception as e:
            ultimo = e
            log(f"   falhou: {msg_curta(e)}")
    raise RuntimeError(
        "Chrome não abre. Feche TODAS as janelas do Chrome deste monitor (ou reinicie o PC) e "
        f"confira se o Chrome está instalado/atualizado. Último erro: {msg_curta(ultimo)}")


# ---- Outlook ----
SEL_ITENS_EMAIL = 'div[role="option"], div[data-convid], div[role="listitem"][aria-label]'
SEL_OUTLOOK_PRONTO = ('div[role="listbox"], div[role="option"], div[data-convid], '
                      'div[role="main"], button[aria-label*="Novo email"], button[aria-label*="New mail"]')


def outlook_pronto(driver):
    """Caixa carregada (mesmo vazia) e fora da tela de login."""
    url = driver.current_url.lower()
    if "login" in url or "microsoftonline" in url or "outlook" not in url:
        return False
    return bool(driver.find_elements(By.CSS_SELECTOR, SEL_OUTLOOK_PRONTO))


def ler_lista(driver):
    itens, vistos = [], set()
    for el in driver.find_elements(By.CSS_SELECTOR, SEL_ITENS_EMAIL):
        try:
            txt = " ".join(f"{el.get_attribute('aria-label') or ''} {el.text or ''}".split())
            cid = el.get_attribute("data-convid") or el.get_attribute("id") or ""
        except StaleElementReferenceException:
            continue
        if not txt or (cid, txt) in vistos:
            continue
        vistos.add((cid, txt))
        itens.append((cid, txt))
    return itens


NOMES_ABAS_CAIXA = ("destaques", "outros", "focused", "other")


def abas_caixa(driver):
    abas = []
    for el in driver.find_elements(By.CSS_SELECTOR, 'button[role="tab"], div[role="tab"]'):
        try:
            if el.is_displayed() and " ".join((el.text or "").split()).lower() in NOMES_ABAS_CAIXA:
                abas.append(el)
        except StaleElementReferenceException:
            pass
    return abas


def ler_caixa(driver):
    """Lê a lista da caixa. Com 'Caixa de Entrada Destaques' ligada, lê Destaques E Outros."""
    abas = abas_caixa(driver)
    if len(abas) < 2:
        return ler_lista(driver)
    itens, vistos = [], set()
    for i in range(len(abas)):
        try:
            aba = abas_caixa(driver)[i]
            clicar(driver, aba)
            time.sleep(2)
        except Exception:
            continue
        for it in ler_lista(driver):
            if it not in vistos:
                vistos.add(it)
                itens.append(it)
    try:
        clicar(driver, abas_caixa(driver)[0])   # volta para Destaques
        time.sleep(1)
    except Exception:
        pass
    return itens


SEL_BUSCA = ['input#topSearchInput', 'input[aria-label="Pesquisar"]', 'input[placeholder="Pesquisar"]',
             'input[aria-label="Search"]', 'input[placeholder="Search"]', 'input[role="combobox"][type="search"]']


def buscar_viradas(driver, termo=None):
    """Faz a busca no Outlook (padrão 'virada', Todas as pastas) e lê os resultados.
    Retorna None se não achou o campo de busca."""
    termo = termo or TERMO_BUSCA
    from selenium.webdriver.common.keys import Keys
    campo = achar_visivel(driver, SEL_BUSCA)
    if campo is None:
        return None
    clicar(driver, campo)
    time.sleep(0.5)
    campo = achar_visivel(driver, SEL_BUSCA) or campo
    campo.send_keys(Keys.CONTROL, "a")
    campo.send_keys(Keys.DELETE)
    campo.send_keys(termo)
    campo.send_keys(Keys.ENTER)
    time.sleep(2)
    # se a busca ficou só na pasta atual, clica "Mostrar resultados de todas as pastas"
    for el in driver.find_elements(By.XPATH, "//*[self::a or self::button or @role='link']"
                                             "[contains(translate(., 'TP', 'tp'), 'todas as pastas')]"):
        try:
            if el.is_displayed():
                clicar(driver, el)
                time.sleep(2)
                break
        except StaleElementReferenceException:
            pass
    # espera os resultados aparecerem
    limite = time.time() + 15
    itens = []
    while time.time() < limite:
        time.sleep(1.5)
        itens = ler_lista(driver)
        if itens:
            break
    return itens


SEL_CORPO_LEITURA = ['div[role="document"]', 'div[aria-label="Corpo da mensagem"]',
                     'div[aria-label="Message body"]', 'div#UniqueMessageBody']


def ler_email_da_lista(driver, deve_conter, preferir=None, espera=20):
    """Clica no item da lista que contém todos os textos de 'deve_conter' (prefere o que tiver 'preferir')
    e devolve o texto do email aberto. None se não achou."""
    alvo = None
    for el in driver.find_elements(By.CSS_SELECTOR, SEL_ITENS_EMAIL):
        try:
            t = " ".join(f"{el.get_attribute('aria-label') or ''} {el.text or ''}".split()).upper()
        except StaleElementReferenceException:
            continue
        if all(c.upper() in t for c in deve_conter):
            if alvo is None or (preferir and preferir.upper() in t):
                alvo = el
    if alvo is None:
        return None
    clicar(driver, alvo)
    limite = time.time() + espera
    texto = ""
    while time.time() < limite:
        time.sleep(1.5)
        partes = []
        for sel in SEL_CORPO_LEITURA:
            for e in driver.find_elements(By.CSS_SELECTOR, sel):
                try:
                    if e.is_displayed():
                        partes.append(driver.execute_script("return arguments[0].innerText;", e) or "")
                except StaleElementReferenceException:
                    pass
        texto = "\n".join(partes)
        if RE_LINHA_ABERTURA.search(texto):
            return texto
    # último recurso: texto da página inteira
    try:
        return driver.execute_script("return document.body.innerText;") or texto
    except Exception:
        return texto


def fechar_busca(driver):
    from selenium.webdriver.common.keys import Keys
    try:
        campo = achar_visivel(driver, SEL_BUSCA)
        if campo:
            campo.send_keys(Keys.ESCAPE)
    except Exception:
        pass


def _dialogo_rotulo(driver):
    for d in driver.find_elements(By.CSS_SELECTOR, 'div[role="dialog"], div[role="alertdialog"]'):
        try:
            if d.is_displayed() and "rótulo" in (d.text or "").lower():
                return d
        except StaleElementReferenceException:
            pass
    return None


def aplicar_rotulo(driver, rotulo=None, log=print, espera=6):
    """Se o Outlook abrir 'Adicionar rótulo de confidencialidade', escolhe o rótulo e clica Enviar.
    Retorna True se tratou a janela, False se ela não apareceu."""
    rotulo = rotulo or ROTULO_CONFIDENCIALIDADE
    limite = time.time() + espera
    dlg = None
    while time.time() < limite and dlg is None:
        dlg = _dialogo_rotulo(driver)
        if dlg is None:
            time.sleep(0.5)
    if dlg is None:
        return False
    log(f"   Outlook pediu rótulo de confidencialidade - escolhendo '{rotulo}'")
    # abre a lista "Selecionar um rótulo"
    abriu = False
    for b in dlg.find_elements(By.XPATH, ".//button | .//*[@role='combobox']"):
        try:
            t = " ".join((b.text or "").split()).lower()
            if b.is_displayed() and ("rótulo" in t or b.get_attribute("role") == "combobox"
                                     or b.get_attribute("aria-haspopup")):
                clicar(driver, b)
                abriu = True
                break
        except StaleElementReferenceException:
            pass
    if not abriu:
        raise RuntimeError("não achei a lista de rótulos na janela do Outlook")
    time.sleep(1)
    # escolhe a opção (a lista abre fora da janela, em outra camada)
    alvo = rotulo.lower()
    escolhido = False
    for _ in range(10):
        for o in driver.find_elements(By.XPATH, "//*[@role='option' or @role='menuitemradio' "
                                                "or @role='menuitem' or @role='menuitemcheckbox']"):
            try:
                if o.is_displayed() and " ".join((o.text or "").split()).lower() == alvo:
                    clicar(driver, o)
                    escolhido = True
                    break
            except StaleElementReferenceException:
                pass
        if escolhido:
            break
        time.sleep(0.5)
    if not escolhido:
        raise RuntimeError(f"opção de rótulo '{rotulo}' não encontrada")
    time.sleep(1)
    # clica Enviar da janela (habilita depois de escolher o rótulo)
    for _ in range(10):
        dlg = _dialogo_rotulo(driver)
        if dlg is None:
            return True
        for b in dlg.find_elements(By.XPATH, ".//button[normalize-space()='Enviar' or .//span[normalize-space()='Enviar']]"):
            try:
                if b.is_displayed() and b.is_enabled() and b.get_attribute("aria-disabled") != "true":
                    clicar(driver, b)
                    log(f"   rótulo '{rotulo}' aplicado, enviando")
                    time.sleep(1)
                    return True
            except StaleElementReferenceException:
                pass
        time.sleep(0.5)
    raise RuntimeError("botão Enviar da janela de rótulo não habilitou")


SELETORES_ENVIAR = [
    (By.CSS_SELECTOR, 'button[aria-label="Enviar"]'),
    (By.CSS_SELECTOR, 'button[aria-label="Send"]'),
    (By.CSS_SELECTOR, 'button[title^="Enviar"]'),
    (By.XPATH, "//button[.//span[normalize-space()='Enviar' or normalize-space()='Send']]"),
]


def achar_botao_enviar(driver):
    for by, sel in SELETORES_ENVIAR:
        for e in driver.find_elements(by, sel):
            try:
                if e.is_displayed() and e.is_enabled():
                    return e
            except StaleElementReferenceException:
                pass
    return None


def textos_dialogos(driver):
    out = []
    for d in driver.find_elements(By.CSS_SELECTOR, 'div[role="dialog"], div[role="alertdialog"]'):
        try:
            if d.is_displayed():
                t = " ".join((d.text or "").split())
                if t:
                    out.append(t[:300])
        except StaleElementReferenceException:
            pass
    return out


def clicar(driver, el):
    try:
        el.click()
    except Exception:
        driver.execute_script("arguments[0].click();", el)


def rascunho_fechou(driver, segundos=6):
    """Enviado de verdade = botão Enviar sumiu e ficou sumido por alguns segundos."""
    sumido = 0
    for _ in range(segundos * 3):
        time.sleep(1)
        if achar_botao_enviar(driver) is None and not textos_dialogos(driver):
            sumido += 1
            if sumido >= segundos:
                return True
        else:
            sumido = 0
    return False


def descartar_rascunho(driver):
    for sel in ('button[aria-label="Descartar"]', 'button[aria-label="Discard"]',
                'button[title^="Descartar"]'):
        els = [e for e in driver.find_elements(By.CSS_SELECTOR, sel) if e.is_displayed()]
        if els:
            clicar(driver, els[0])
            time.sleep(2)
            for e in driver.find_elements(By.XPATH,
                    "//div[@role='dialog' or @role='alertdialog']//button[normalize-space()='Descartar' "
                    "or normalize-space()='OK' or normalize-space()='Discard']"):
                if e.is_displayed():
                    clicar(driver, e)
                    break
            time.sleep(2)
            return True
    return False


SEL_ASSUNTO = ['input[aria-label="Adicionar um assunto"]', 'input[placeholder="Adicionar um assunto"]',
               'input[aria-label="Assunto"]', 'input[aria-label="Add a subject"]',
               'input[placeholder="Add a subject"]']
SEL_CORPO = ['div[role="textbox"][aria-label="Corpo da mensagem"]', 'div[aria-label="Corpo da mensagem"]',
             'div[role="textbox"][aria-label="Message body"]', 'div[contenteditable="true"][aria-label*="orpo"]',
             'div[contenteditable="true"][aria-label*="body"]']


def achar_visivel(driver, seletores):
    for sel in seletores:
        for e in driver.find_elements(By.CSS_SELECTOR, sel):
            try:
                if e.is_displayed():
                    return e
            except StaleElementReferenceException:
                pass
    return None


def achar_corpo(driver):
    e = achar_visivel(driver, SEL_CORPO)
    if e:
        return e
    # fallback: maior campo editável que não seja Para/Cc/Cco
    melhor, alt = None, 0
    for c in driver.find_elements(By.CSS_SELECTOR, 'div[contenteditable="true"]'):
        try:
            lbl = (c.get_attribute("aria-label") or "").lower()
            if lbl in ("para", "cc", "cco", "to", "bcc") or not c.is_displayed():
                continue
            h = c.size.get("height", 0)
            if h > alt:
                melhor, alt = c, h
        except StaleElementReferenceException:
            pass
    return melhor


def origem(url):
    u = urllib.parse.urlsplit(url)
    return f"{u.scheme}://{u.netloc}" if u.netloc else "https://outlook.office.com"


def enviar_email(driver, para, cc, assunto, corpo, log=print):
    from selenium.webdriver.common.keys import Keys
    from selenium.webdriver.common.action_chains import ActionChains

    params = {"to": para}
    if cc:
        params["cc"] = ";".join(cc)
    url = origem(driver.current_url) + OUTLOOK_COMPOSE_PATH + "?" + \
        urllib.parse.urlencode(params, quote_via=urllib.parse.quote)

    original = driver.current_window_handle
    driver.switch_to.new_window("tab")
    enviado = False
    try:
        log("   abrindo novo email...")
        driver.get(url)
        limite = time.time() + 90
        while achar_botao_enviar(driver) is None or achar_visivel(driver, SEL_ASSUNTO) is None:
            if time.time() > limite:
                raise RuntimeError("Tela de novo email não carregou")
            time.sleep(1)

        # destinatário (vem pelo link; confere)
        local = para.split("@")[0].lower()
        limite = time.time() + 20
        while True:
            fonte = driver.page_source.lower()
            if local in fonte or local.replace(".", " ") in fonte:
                break
            if time.time() > limite:
                raise RuntimeError("Destinatário não apareceu no rascunho")
            time.sleep(1)
        time.sleep(2)

        log("   destinatário OK")
        # assunto: digitado (o Outlook novo ignora o assunto do link)
        campo = achar_visivel(driver, SEL_ASSUNTO)
        clicar(driver, campo)
        campo.send_keys(Keys.CONTROL, "a")
        campo.send_keys(assunto)
        time.sleep(1)
        if (campo.get_attribute("value") or "").strip() != assunto.strip():
            raise RuntimeError("Assunto não ficou preenchido")

        log("   assunto preenchido")
        # corpo: digitado no começo, acima da assinatura automática
        caixa = achar_corpo(driver)
        if caixa is None:
            raise RuntimeError("Corpo do email não encontrado")
        clicar(driver, caixa)
        time.sleep(0.5)
        acoes = ActionChains(driver)
        acoes.key_down(Keys.CONTROL).send_keys(Keys.HOME).key_up(Keys.CONTROL)
        for i, linha in enumerate(corpo.split("\n")):
            if i:
                acoes.send_keys(Keys.ENTER)
            if linha:
                acoes.send_keys(linha)
        acoes.send_keys(Keys.ENTER).perform()
        time.sleep(1.5)
        marcador = corpo.split("\n")[2][:30] if len(corpo.split("\n")) > 2 else corpo[:30]
        if marcador.lower() not in (caixa.text or "").lower():
            raise RuntimeError("Texto não entrou no corpo do email")

        log("   corpo preenchido")
        # enviar
        btn = achar_botao_enviar(driver)
        if btn is None:
            raise RuntimeError("Botão Enviar sumiu antes do clique")
        clicar(driver, btn)
        aplicar_rotulo(driver, log=log)
        log("   clique em Enviar feito, aguardando o Outlook confirmar...")
        if rascunho_fechou(driver):
            enviado = True
            time.sleep(4)
            return

        dialogos = textos_dialogos(driver)
        try:
            driver.save_screenshot(str(BASE / "falha_envio.png"))
        except Exception:
            pass
        if dialogos:
            log(f"   janela do Outlook bloqueou o envio: {' | '.join(dialogos)}")
        ActionChains(driver).key_down(Keys.CONTROL).send_keys(Keys.ENTER).key_up(Keys.CONTROL).perform()
        log("   tentando Ctrl+Enter...")
        aplicar_rotulo(driver, log=log)
        if rascunho_fechou(driver):
            enviado = True
            time.sleep(4)
            return
        detalhe = " | ".join(textos_dialogos(driver) or dialogos) or "sem mensagem na tela"
        raise RuntimeError(f"Outlook não enviou ({detalhe}) - print em falha_envio.png")
    except Exception:
        try:
            driver.save_screenshot(str(BASE / "falha_envio.png"))
        except Exception:
            pass
        raise
    finally:
        try:
            if not enviado and achar_botao_enviar(driver) is not None:
                if descartar_rascunho(driver):
                    log("   rascunho descartado (para não acumular rascunhos)")
        except Exception:
            pass
        try:
            driver.close()
        finally:
            driver.switch_to.window(original)


# ---- SharePoint ----
JS_BAIXAR = """
var done = arguments[arguments.length - 1];
fetch(arguments[0], {credentials: 'include', cache: 'no-store'}).then(function (r) {
  if (!r.ok) { done('ERRO:' + r.status); return null; }
  return r.arrayBuffer();
}).then(function (b) {
  if (!b) return;
  var u = new Uint8Array(b), s = '', CH = 0x8000;
  for (var i = 0; i < u.length; i += CH) { s += String.fromCharCode.apply(null, u.subarray(i, i + CH)); }
  done(btoa(s));
}).catch(function (e) { done('ERRO:' + e); });
"""


def sharepoint_pronto(driver):
    url = driver.current_url.lower()
    return "sharepoint.com" in url and "login" not in url and "microsoftonline" not in url


JS_TEXTO = """
var done = arguments[arguments.length - 1];
fetch(arguments[0], {credentials: 'include', cache: 'no-store',
                     headers: {'Accept': 'application/json;odata=nometadata'}})
  .then(function (r) { return r.ok ? r.text() : 'ERRO:' + r.status; })
  .then(done).catch(function (e) { done('ERRO:' + e); });
"""


def _alias(caminho):
    return urllib.parse.quote("'" + caminho.replace("'", "''") + "'", safe="/")


def sp_baixar(driver, url):
    import base64
    driver.set_script_timeout(90)
    r = driver.execute_async_script(JS_BAIXAR, url)
    if not r or str(r).startswith("ERRO:"):
        raise RuntimeError(f"download falhou ({r})")
    dados = base64.b64decode(r)
    if dados[:2] != b"PK":
        raise RuntimeError("SharePoint não devolveu o Excel (sessão expirou? faça login na aba do SharePoint)")
    return dados


def sp_listar(driver, pasta):
    """-> (arquivos [(nome, caminho, modificado)], subpastas [(nome, caminho)])"""
    url = (f"{SP_SITE}/_api/web/GetFolderByServerRelativePath(decodedurl=@a1)"
           f"?@a1={_alias(pasta)}&$expand=Files,Folders"
           f"&$select=Files/Name,Files/ServerRelativeUrl,Files/TimeLastModified,Files/UniqueId,"
           f"Folders/Name,Folders/ServerRelativeUrl")
    driver.set_script_timeout(60)
    r = driver.execute_async_script(JS_TEXTO, url)
    if not r or str(r).startswith("ERRO:"):
        raise RuntimeError(f"não consegui listar a pasta ({r})")
    d = json.loads(r)
    arqs = [(f["Name"], f["ServerRelativeUrl"], f.get("TimeLastModified", ""), f.get("UniqueId", ""))
            for f in d.get("Files", [])]
    pastas = [(f["Name"], f["ServerRelativeUrl"]) for f in d.get("Folders", []) if f["Name"] != "Forms"]
    return arqs, pastas


def ordenar_arquivos_passagem(arquivos, agora=None):
    """Candidatos em ordem de preferência: .xlsx com PASSAGEM no nome; nome com o mês atual;
    sem nome de outro mês; modificado por último."""
    agora = agora or datetime.now()
    mes = MESES[agora.month - 1]
    sem_acento = mes.replace("Ç", "C")
    outros = [m_ for m_ in MESES if m_ != mes] + [m_.replace("Ç", "C") for m_ in MESES if m_ != mes]
    cand = [a for a in arquivos
            if a[0].lower().endswith((".xlsx", ".xlsm")) and not a[0].startswith("~$")]

    def nota(a):
        nome = a[0].upper()
        do_mes = mes in nome or sem_acento in nome or f"{agora.month:02d}.{agora.year}" in nome \
            or f"{agora.month:02d}-{agora.year}" in nome or f"{agora.month:02d}_{agora.year}" in nome
        outro_mes = any(o in nome for o in outros)
        caminho = a[1].upper()
        pasta_mes = f"/{agora.month:02d} -" in caminho or f"/{agora.month:02d}-" in caminho \
            or f"/{mes}/" in caminho or f"/{sem_acento}/" in caminho
        return (NOME_PLANILHA in nome, "DESVIO" not in nome, "PASSAGEM" in nome,
                pasta_mes or do_mes, not outro_mes, a[2])
    return sorted(cand, key=nota, reverse=True)


def escolher_arquivo_passagem(arquivos, agora=None):
    lista = ordenar_arquivos_passagem(arquivos, agora)
    return lista[0] if lista else None


def passagem_do_mes(conteudo, agora=None):
    """True se a aba de hoje tem a data de hoje (célula 'Data'), False se tem data de outro mês,
    None se não achou data."""
    import io
    import openpyxl
    agora = agora or datetime.now()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    except Exception:
        return False
    nome = str(agora.day)
    if nome not in wb.sheetnames:
        return False
    for row in wb[nome].iter_rows(min_row=1, max_row=4, max_col=30, values_only=True):
        for v in row:
            if isinstance(v, datetime):
                return (v.month, v.year) == (agora.month, agora.year)
            m_ = re.search(r"\b\d{1,2}/(\d{1,2})/(\d{4})\b", str(v or ""))
            if m_:
                return (int(m_.group(1)), int(m_.group(2))) == (agora.month, agora.year)
    return None


def baixar_passagem(driver, log=print):
    """Entra na pasta do ano -> pasta do mês ("10 - Outubro") e baixa a planilha A1 e A2.
    -> (bytes, nome, UniqueId do arquivo)"""
    try:
        agora = datetime.now()
        arqs, pastas = sp_listar(driver, f"{SP_RAIZ}/{agora.year}")
        do_mes = [c for n, c in pastas if n.strip().startswith(f"{agora.month:02d}")]
        if do_mes:
            a2, _p = sp_listar(driver, do_mes[0])
            arqs += a2
        else:
            log(f"Passagem: pasta '{agora.month:02d} - ...' não encontrada em {agora.year} - procurando nas outras")
            for _, caminho in pastas:
                try:
                    a2, _p = sp_listar(driver, caminho)
                    arqs += a2
                except Exception:
                    pass
        # baixa até 3 candidatos e fica com o que tem a data de HOJE na aba do dia
        reserva = None
        for nome, caminho, _mod, *resto in ordenar_arquivos_passagem(arqs)[:3]:
            uid = resto[0] if resto else ""
            url = (f"{SP_SITE}/_api/web/GetFileByServerRelativePath(decodedurl=@a1)/$value"
                   f"?@a1={_alias(caminho)}")
            try:
                dados = sp_baixar(driver, url)
            except Exception as e:
                log(f"Passagem: não baixou '{nome}': {msg_curta(e)}")
                continue
            ok = passagem_do_mes(dados)
            if ok:
                return dados, nome, uid
            if ok is None and reserva is None:
                reserva = (dados, nome, uid)
        if reserva:
            return reserva
        log("Passagem: nenhuma planilha da pasta tem a data de hoje - usando o arquivo do link")
    except Exception as e:
        log(f"Passagem: {msg_curta(e)} - usando o arquivo padrão")
    return (sp_baixar(driver, SP_DOWNLOAD + f"&_={int(time.time())}"),
            "Passagem de Turno (arquivo padrão).xlsx", SP_ARQUIVO_ID)


def url_excel_no_dia(uid, dia=None):
    """Link do Excel online da passagem já abrindo na aba do dia (ex.: aba '21')."""
    dia = dia or datetime.now()
    uid = uid.strip("{}")
    celula = urllib.parse.quote(f"'{dia.day}'!A1", safe="!")
    return (f"{SP_SITE}/_layouts/15/Doc.aspx?sourcedoc=%7B{uid}%7D&action=default"
            f"&activeCell={celula}")


# ------------------------------ monitor -----------------------------------
class Monitor(threading.Thread):
    def __init__(self, fila, modo_teste=MODO_TESTE):
        super().__init__(daemon=True)
        self.modo_teste = modo_teste      # alterado pela tela (caixa "Modo teste")
        self.fila = fila
        self.parar = threading.Event()
        self.driver = None
        self.aba_outlook = None
        self.aba_sp = None
        self.passagem = {}
        self.passagem_quando = None
        self.estado = None
        self.pedir_diag = threading.Event()
        self.falhas = {}
        self.pedido_manual = None
        self.pendentes = {}   # virada sem tanque Base -> hora da última tentativa
        self.ligado = None
        self.ultima_abertura = 0.0
        self.ultimo_semanal = 0.0
        self._trocas_iniciado = False
        self._laudos_iniciado = False
        self._avisos_abertura = set()

    def log(self, msg):
        self.fila.put(("log", msg))

    def status(self, msg):
        self.fila.put(("status", msg))

    def run(self):
        while not self.parar.is_set():
            try:
                self.ciclo()
            except Exception as e:
                msg = msg_curta(e)
                if "invalid session" in msg.lower() or "disconnected" in msg.lower() \
                        or "no such window" in msg.lower():
                    self.log("Chrome foi fechado ou travou - reabrindo em 30 s.")
                else:
                    self.log(f"ERRO: {msg}")
                    self.log(traceback.format_exc(limit=1).split("Stacktrace:")[0].strip())
                self.status("Erro - reiniciando navegador em 30 s")
                self.fechar_driver()
                self.parar.wait(30)
        self.fechar_driver()
        self.status("Parado")

    def fechar_driver(self):
        if self.driver:
            try:
                self.driver.quit()
            except Exception:
                pass
            self.driver = None

    def passo(self, msg):
        self.status(msg)
        self.log(msg)

    def esperar(self, condicao, minutos, msg_status):
        """True = condição atendida; False = tempo esgotado ou parado."""
        self.passo(msg_status)
        inicio = time.time()
        prox_aviso = inicio + 60
        while True:
            try:
                if condicao():
                    return True
            except Exception:
                pass
            if self.parar.is_set() or time.time() - inicio > minutos * 60:
                return False
            if self.pedir_diag.is_set():
                self.pedir_diag.clear()
                try:
                    arq = BASE / "diagnostico_espera.png"
                    self.driver.save_screenshot(str(arq))
                    self.log(f"DIAGNÓSTICO (esperando): {self.driver.title} | {self.driver.current_url} "
                             f"| print salvo: {arq.name}")
                except Exception as e:
                    self.log(f"DIAGNÓSTICO: erro {e}")
            if time.time() > prox_aviso:
                prox_aviso += 60
                try:
                    url = self.driver.current_url
                except Exception:
                    url = "?"
                self.log(f"  ...ainda aguardando ({int((time.time() - inicio) / 60)} min) - página: {url}")
            time.sleep(3)

    def ciclo(self):
        self.passo("Abrindo Chrome...")
        self.driver = criar_driver(log=self.log)
        self.log("Chrome aberto.")

        # aba 1: Outlook
        self.passo("Abrindo Outlook...")
        self.driver.get(OUTLOOK_INBOX)
        self.aba_outlook = self.driver.current_window_handle
        if not self.esperar(lambda: outlook_pronto(self.driver), 15,
                            "Aguardando Outlook (faça login se pedir)"):
            if self.parar.is_set():
                return
            raise RuntimeError("Outlook não carregou em 15 min")
        self.log("Outlook OK.")

        # aba 2: SharePoint (passagem de turno)
        self.passo("Abrindo SharePoint (Passagem de Turno)...")
        self.driver.switch_to.new_window("tab")
        self.aba_sp = self.driver.current_window_handle
        try:
            self.driver.get(SP_PASTA_URL)
        except Exception as e:
            self.log(f"SharePoint demorou para abrir: {msg_curta(e)}")
        if self.esperar(lambda: sharepoint_pronto(self.driver), 5,
                        "Aguardando SharePoint (faça login se pedir)"):
            self.atualizar_passagem()
        else:
            if self.parar.is_set():
                return
            self.log("SharePoint não abriu - vou usar a reserva da tela e tentar de novo a cada 10 min.")
        self.driver.switch_to.window(self.aba_outlook)

        # Ao iniciar: TUDO que já está no Outlook é marcado como visto (não envia).
        # Só é enviado email de virada que chegar DEPOIS que o robô foi ligado.
        self.estado = self.estado or carregar_estado() or {"chaves": {}, "historico": []}
        self.estado.setdefault("enviadas", {})
        for reg in self.estado.get("historico", []):        # compatibilidade com versões anteriores
            if str(reg.get("status", "")).startswith("Enviado"):
                try:
                    q = datetime.strptime(reg["quando"], "%d/%m/%Y %H:%M")
                    kv = f"VIRADA|{int(reg['tanque'])}|{reg['hora']}"
                    self.estado["enviadas"].setdefault(kv, q.isoformat())
                except Exception:
                    pass
        if self.ligado is None:
            self.ligado = datetime.now()      # hora em que o robô foi ligado (não muda se o Chrome reabrir)
        ligado = self.ligado
        agora = datetime.now().isoformat()
        itens = self.ler_viradas()
        existentes = 0
        for cid, txt in itens:
            v = extrair_virada(txt)
            if not v:
                continue
            self.estado["chaves"][self.chave(cid, v)] = agora
            existentes += 1
            chegou = hora_do_item(txt)
            if chegou is not None:            # virada de hoje: serve para saber o último tanque do pool
                self.registrar_pool(v["produto"], v["tanque"], data_virada(v, chegou))
        self.verificar_trocas_base(inicial=not self._trocas_iniciado)
        self.verificar_laudos_hidratado(inicial=not self._laudos_iniciado)
        self._laudos_iniciado = True
        self._trocas_iniciado = True
        salvar_estado(self.estado)
        self.log(f"Robô ligado {ligado:%d/%m %H:%M}. Busca '{TERMO_BUSCA}': {existentes} virada(s) já existente(s) "
                 f"ignorada(s) - só envio email de virada que chegar a partir de agora.")

        self.status("Monitorando")
        self.log("Monitorando Outlook.")
        ultimo_f5 = ultimo_resumo = ultimo_passagem = time.time()

        while not self.parar.is_set():
            if self.pedir_diag.is_set():
                self.pedir_diag.clear()
                self.diagnostico()

            if self.pedido_manual:
                v, self.pedido_manual = self.pedido_manual, None
                self.log(f"Envio MANUAL pedido na tela: TQ {v['tanque']} {v['produto']} {v['hora']}")
                self.processar(v, self.chave(f"manual-{time.time():.0f}", v))

            if time.time() - ultimo_resumo > 300:
                ultimo_resumo = time.time()
                self.log(f"Varredura OK - busca '{TERMO_BUSCA}' rodando, {len(self.pendentes)} pendente(s).")

            if time.time() - ultimo_passagem > INTERVALO_PASSAGEM_S:
                self.atualizar_passagem()
                self.verificar_trocas_base()
                self.verificar_laudos_hidratado()
                ultimo_passagem = time.time()

            if time.time() - self.ultimo_semanal > RETENTAR_PENDENTE_S:
                self.ultimo_semanal = time.time()
                self.verificar_semanal_s10b30()

            intervalo_ab = INTERVALO_ABERTURA_S if self.na_janela_abertura() else (
                None if self.estado.get("aberturas", {}).get(datetime.now().strftime("%d/%m/%Y"))
                else INTERVALO_ABERTURA_FORA_S)
            if intervalo_ab and time.time() - self.ultima_abertura > intervalo_ab:
                self.ultima_abertura = time.time()
                try:
                    self.verificar_abertura()
                except Exception as e:
                    if "invalid session" in str(e).lower() or "no such window" in str(e).lower():
                        raise
                    self.log(f"Abertura: erro {msg_curta(e)}")

            self.verificar_aviso_sem_abertura()

            if time.time() - ultimo_f5 > RECARREGAR_A_CADA_S:
                self.recarregar_abas()
                ultimo_f5 = time.time()

            vistos_ciclo = set()
            for cid, txt in self.ler_viradas():
                v = extrair_virada(txt)
                if not v:
                    continue
                k = self.chave(cid, v)
                if k in self.estado["chaves"] or self.chave_virada(v) in vistos_ciclo:
                    continue
                vistos_ciclo.add(self.chave_virada(v))
                # só email que CHEGOU depois que o robô foi ligado (pelo horário da lista do Outlook)
                chegou = hora_do_item(txt)
                if chegou is None or chegou < self.ligado - timedelta(minutes=2):
                    self.estado["chaves"][k] = datetime.now().isoformat()
                    salvar_estado(self.estado)
                    quando = chegou.strftime("%H:%M") if chegou else "dia anterior"
                    self.log(f"TQ {v['tanque']} {v['produto']} {v['hora']}: email antigo ({quando}, "
                             f"antes do robô ligar às {self.ligado:%H:%M}) - ignorado")
                    continue
                # resposta/encaminhamento hoje de uma virada velha (ex.: ENC às 02:05 da virada de 21h00 de ontem)
                dv = data_virada(v, chegou)
                if abs(chegou - dv) > timedelta(hours=HORAS_EMAIL_X_VIRADA):
                    self.estado["chaves"][k] = datetime.now().isoformat()
                    salvar_estado(self.estado)
                    self.log(f"TQ {v['tanque']} {v['produto']} {v['hora']}: virada de {dv:%d/%m %H:%M} respondida/"
                             f"encaminhada às {chegou:%H:%M} - virada antiga, ignorado")
                    continue
                if self.virada_recente(v):
                    # mesma virada já pedida (RE:, encaminhado, outra conversa)
                    self.estado["chaves"][k] = datetime.now().isoformat()
                    salvar_estado(self.estado)
                    self.log(f"TQ {v['tanque']} {v['produto']} {v['hora']}: já pedido antes - ignorado (RE/encaminhado)")
                    continue
                if self.falhas.get(k, 0) >= MAX_TENTATIVAS:
                    continue
                if k in self.pendentes and time.time() - self.pendentes[k] < RETENTAR_PENDENTE_S:
                    continue
                self.processar(v, k)

            self.fila.put(("tick", datetime.now().strftime("%H:%M:%S")))
            self.parar.wait(INTERVALO_S)

    def diagnostico(self):
        self.log("===== DIAGNÓSTICO =====")
        for nome, aba in (("Outlook", self.aba_outlook),):
            if not aba:
                self.log(f"{nome}: aba não aberta")
                continue
            try:
                self.driver.switch_to.window(aba)
                d = self.driver
                self.log(f"{nome}: {d.title} | {d.current_url}")
                for sel in ('div[role="option"]', 'div[data-convid]', 'div[role="listitem"]',
                            'div[role="listbox"]'):
                    self.log(f"   {sel}: {len(d.find_elements(By.CSS_SELECTOR, sel))}")
                arq = BASE / f"diagnostico_{nome.lower()}.png"
                d.save_screenshot(str(arq))
                self.log(f"   print salvo: {arq.name}")
            except Exception as e:
                self.log(f"{nome}: erro no diagnóstico: {e}")
        try:
            self.driver.switch_to.window(self.aba_outlook)
            itens = ler_lista(self.driver)
            for _, t in itens[:5]:
                self.log(f"   email: {t[:150]}")
        except Exception as e:
            self.log(f"   erro lendo emails: {e}")
        self.log("=======================")

    def recarregar_abas(self):
        self.driver.switch_to.window(self.aba_outlook)
        self.driver.refresh()
        time.sleep(8)

    @staticmethod
    def chave(cid, v):
        return f"{cid}|{v['produto']}|{int(v['tanque'])}|{v['hora']}"

    @staticmethod
    def chave_virada(v):
        # mesma virada = mesmo tanque + mesmo horário (produto pode vir escrito diferente: S500 / OD S500)
        return f"VIRADA|{int(v['tanque'])}|{v['hora']}"

    def ler_viradas(self):
        """Busca 'virada' no Outlook; se não achar o campo de busca, lê a caixa de entrada."""
        itens = buscar_viradas(self.driver)
        if itens is None:
            if not getattr(self, "_avisou_busca", False):
                self.log("Campo de busca do Outlook não encontrado - lendo a caixa de entrada")
                self._avisou_busca = True
            return ler_caixa(self.driver)
        return itens

    def atualizar_passagem(self):
        """Baixa a planilha pela aba do SharePoint e volta para o Outlook."""
        try:
            self.driver.switch_to.window(self.aba_sp)
            if not sharepoint_pronto(self.driver):
                self.driver.get(SP_PASTA_URL)
                self._excel_aberto = None
                time.sleep(5)
            dados, nome, uid = baixar_passagem(self.driver, log=self.log)
            self.passagem_bytes = dados
            p = ler_passagem(dados)
            if nome != getattr(self, "passagem_arquivo", None):
                self.log(f"Passagem: lendo '{nome}'")
                self.passagem_arquivo = nome
            if not p:
                self.log("Passagem: não achei B100/EA na aba do dia - usando a reserva da tela")
            self.passagem, self.passagem_quando = p, datetime.now()
            self.fila.put(("passagem", (self.passagem_quando.strftime("%H:%M"), p)))
            # deixa a aba do SharePoint com a planilha aberta na aba do DIA (troca sozinho à meia-noite)
            alvo = (uid, datetime.now().strftime("%d/%m/%Y"))
            if uid and getattr(self, "_excel_aberto", None) != alvo:
                self.driver.get(url_excel_no_dia(uid))
                self._excel_aberto = alvo
                self.log(f"Passagem: planilha aberta no Excel online na aba do dia {datetime.now().day}")
                time.sleep(3)
        except Exception as e:
            self.log(f"Passagem: erro na leitura: {msg_curta(e)}")
        finally:
            try:
                self.driver.switch_to.window(self.aba_outlook)
            except Exception:
                pass

    def virada_recente(self, v):
        """Essa virada (tanque + horário) já foi ENVIADA nas últimas horas?"""
        t = self.estado.setdefault("enviadas", {}).get(self.chave_virada(v))
        return bool(t) and datetime.fromisoformat(t) > datetime.now() - timedelta(hours=HORAS_MESMA_VIRADA)

    # ---------- envio genérico ----------
    def enviar_pedido(self, assunto, corpo, reg, k):
        """Envia (Laboratório ou teste). True se enviou."""
        try:
            para, cc = (EMAIL_TESTE, []) if self.modo_teste else (EMAIL_LABORATORIO, EMAIL_CC)
            reg["para"] = para
            self.log(f"Destinatário: {para}" + ("  [MODO TESTE]" if self.modo_teste else ""))
            enviar_email(self.driver, para, cc, assunto, corpo, log=self.log)
        except Exception as e:
            self.falhas[k] = self.falhas.get(k, 0) + 1
            reg["status"] = (f"DESISTIU após {MAX_TENTATIVAS} tentativas: {e}" if self.falhas[k] >= MAX_TENTATIVAS
                             else f"FALHA: {e}")
            self.log(f"Falha no envio ({self.falhas[k]}ª vez): {e}")
            self.fila.put(("enviado", reg))
            return False
        reg["status"] = f"Enviado (TESTE -> {para})" if self.modo_teste else "Enviado"
        self.estado["historico"].append(reg)
        self.fila.put(("enviado", reg))
        return True

    def registrar_pool(self, produto, tanque, quando):
        """Guarda o último tanque do pool por produto (S10 / S500 / GAS)."""
        pp = produto_pool(produto)
        if not pp or not tanque:
            return
        ult = self.estado.setdefault("ultimo_pool", {})
        atual = ult.get(pp)
        if atual and atual.get("quando", "") > quando.isoformat():
            return
        ult[pp] = {"tanque": f"{int(tanque):02d}", "quando": quando.isoformat()}

    # ---------- troca de tanque Base (B100 / Anidro) ----------
    # ---------- hidratado: linha nova na seção LAUDOS ----------
    def verificar_laudos_hidratado(self, inicial=False):
        dados = getattr(self, "passagem_bytes", None)
        if self.estado is None or not dados:
            return
        try:
            linhas = ler_laudos_passagem(dados)
        except Exception as e:
            self.log(f"Laudos (hidratado): erro lendo a passagem: {msg_curta(e)}")
            return
        resumo = " | ".join(f"aba {ln['aba']} lin {ln['linha']}: TQ {ln['tanque']} amostra='{ln['amostra']}' "
                            f"emitido='{ln['emitido']}'" for ln in linhas) or "nenhum tanque de hidratado lançado"
        if resumo != getattr(self, "_ultimo_resumo_laudos", None):
            self._ultimo_resumo_laudos = resumo
            self.log(f"Seção LAUDOS lida: {resumo}")
        vistos = self.estado.setdefault("laudos_vistos", {})
        enviadas = self.estado.setdefault("enviadas", {})
        for ln in linhas:
            if not ln["amostra"] or ln["emitido"]:
                continue                       # amostra ainda não retirada / laudo já emitido
            k = f"LAUDO|{ln['aba']}|{ln['linha']}|{ln['tanque']}"
            if k in vistos or k in enviadas:
                continue
            if inicial:
                vistos[k] = datetime.now().isoformat()   # já estava lançado quando o robô ligou
                self.log(f"LAUDOS: TQ {ln['tanque']} já estava lançado quando o robô ligou - não pedido "
                         f"(use 'Enviar agora' se precisar)")
                continue
            if self.falhas.get(k, 0) >= MAX_TENTATIVAS:
                continue
            t = ln["tanque"]
            self.log(f"LAUDOS (passagem aba {ln['aba']}, linha {ln['linha']}): amostra do TQ {t} hidratado "
                     f"retirada ({ln['amostra']}) - pedindo laudo")
            assunto, corpo = montar_email_livre(
                f"Solicitação de laudo - TQ {t} (Base) - ETANOL HIDRATADO",
                f"Foi retirada amostra do TQ {t} (ETANOL HIDRATADO) e a amostra já está na caixinha.",
                [f"TQ {t} (Base) - ETANOL HIDRATADO"])
            reg = {"quando": datetime.now().strftime("%d/%m/%Y %H:%M"), "tanque": "-",
                   "produto": "AMOSTRA ETANOL HIDRATADO", "hora": "", "tq_base": t}
            if self.enviar_pedido(assunto, corpo, reg, k):
                enviadas[k] = datetime.now().isoformat()
        salvar_estado(self.estado)

    def verificar_trocas_base(self, inicial=False):
        if self.estado is None or not self.passagem:
            return
        guardado = self.estado.setdefault("tanques_base", {})
        for prod in PRODUTOS_TROCA:
            atual = tanques_passagem(prod, self.passagem)
            if not atual:
                continue                      # célula vazia / não lida: não é troca
            anterior = guardado.get(prod)
            if inicial or anterior is None:
                guardado[prod] = atual
                continue
            novos = [t for t in atual if t not in anterior]
            if not novos:
                guardado[prod] = atual
                continue
            info = self.passagem.get(prod, {})
            hoje = datetime.now().strftime("%d/%m/%Y")
            enviados_ok = True
            for t in novos:
                k = f"TROCA|{prod}|{t}|{hoje}"
                if k in self.estado.setdefault("enviadas", {}) or self.falhas.get(k, 0) >= MAX_TENTATIVAS:
                    continue
                self.log(f"TROCA de tanque {prod}: {' + '.join(anterior)} -> {' + '.join(atual)} "
                         f"(passagem aba {info.get('aba')}, {info.get('turno')}º turno) - pedindo laudo do TQ {t}")
                self.garantir_abertura_hoje()
                pools = []
                for pp in POOL_DO_COMPONENTE.get(prod, ()):
                    up = self.estado.get("ultimo_pool", {}).get(pp)
                    if up:
                        pools.append((pp, up["tanque"]))
                    else:
                        self.log(f"   último tanque do pool de {pp} desconhecido - vai sem ele")
                linhas = [f"TQ {t} (Base, em expedição) - {prod}"] + [f"TQ {tq} (Pool) - {pp}" for pp, tq in pools]
                pool_txt = "".join(f" + TQ {tq} (Pool {pp})" for pp, tq in pools)
                assunto, corpo = montar_email_livre(
                    f"Solicitação de laudo - TQ {t} (Base){pool_txt} - troca {prod}",
                    f"Houve troca do tanque em expedição de {prod}: TQ {' + '.join(anterior)} -> TQ {t}.", linhas)
                reg = {"quando": datetime.now().strftime("%d/%m/%Y %H:%M"),
                       "tanque": " + ".join(tq for _, tq in pools) or "-",
                       "produto": f"TROCA {prod}", "hora": "", "tq_base": t}
                if self.enviar_pedido(assunto, corpo, reg, k):
                    self.estado["enviadas"][k] = datetime.now().isoformat()
                else:
                    enviados_ok = False
            if enviados_ok:
                guardado[prod] = atual        # se falhou, tenta de novo na próxima leitura
        salvar_estado(self.estado)

    # ---------- pedido semanal S10 B30 (segunda 00:30) ----------
    def verificar_semanal_s10b30(self, agora=None):
        agora = agora or datetime.now()
        if agora.weekday() != SEMANAL_DIA or not (SEMANAL_HORA <= (agora.hour, agora.minute) < SEMANAL_ATE):
            return
        k = f"SEMANAL|S10B30|{agora:%d/%m/%Y}"
        if k in self.estado.setdefault("enviadas", {}) or self.falhas.get(k, 0) >= MAX_TENTATIVAS:
            return
        self.atualizar_passagem()
        self.garantir_abertura_hoje()
        _p, tanques_rz, origem_tq = escolher_tanques("OD S10", self.passagem)
        pool = self.estado.get("ultimo_pool", {}).get("S10")
        faltando = ([] if pool else ["último tanque de S10 do pool"]) + ([] if tanques_rz else ["tanque de B100 expedindo"])
        if faltando:
            if getattr(self, "_aviso_semanal", None) != k:
                self._aviso_semanal = k
                self.log(f"AGUARDANDO: laudo semanal S10 B30 - falta {' e '.join(faltando)}. "
                         f"Tento de novo a cada {RETENTAR_PENDENTE_S // 60} min.")
                self.fila.put(("alerta", "Laudo semanal S10 B30 aguardando: falta " + " e ".join(faltando)))
            return
        rz = [numero_tanque(n) for n in tanques_rz]
        assunto, corpo = montar_email_livre(
            f"Solicitação de laudo semanal - S10 B30 - TQ {pool['tanque']} (Pool) + TQ {' + TQ '.join(rz)} (Base)",
            "Solicitação semanal (segunda-feira) do laudo do S10 B30.",
            [f"TQ {pool['tanque']} (Pool) - S10"] + [f"TQ {t} (Base, em expedição) - BIODIESEL B100" for t in rz])
        self.log(f"Laudo semanal S10 B30: pool S10 TQ {pool['tanque']} + B100 TQ {'+'.join(rz)} ({origem_tq}) -> enviando")
        reg = {"quando": agora.strftime("%d/%m/%Y %H:%M"), "tanque": pool["tanque"],
               "produto": "SEMANAL S10 B30", "hora": "", "tq_base": " + ".join(rz)}
        if self.enviar_pedido(assunto, corpo, reg, k):
            self.estado["enviadas"][k] = datetime.now().isoformat()
            salvar_estado(self.estado)

    # ---------- email ABERTURA DE TANQUES ----------
    @staticmethod
    def na_janela_abertura(agora=None):
        agora = agora or datetime.now()
        hm = (agora.hour, agora.minute)
        ini, fim = JANELA_ABERTURA
        return hm >= ini or hm < fim

    def garantir_abertura_hoje(self):
        """Antes de usar o 'último tanque do pool', tenta ler a abertura de hoje se ainda não leu."""
        hoje_txt = datetime.now().strftime("%d/%m/%Y")
        if self.estado.get("aberturas", {}).get(hoje_txt) or time.time() - self.ultima_abertura < 60:
            return
        self.ultima_abertura = time.time()
        try:
            self.verificar_abertura()
        except Exception as e:
            if "invalid session" in str(e).lower() or "no such window" in str(e).lower():
                raise
            self.log(f"Abertura: erro {msg_curta(e)}")
        if not self.estado.get("aberturas", {}).get(hoje_txt):
            self.log(f"   abertura de {hoje_txt} ainda não lida - usando o último tanque do pool conhecido")

    def verificar_aviso_sem_abertura(self, agora=None):
        """Às 00:15, se ainda não leu a abertura do dia, avisa o responsável (uma vez por dia)."""
        agora = agora or datetime.now()
        hora = agora.replace(hour=AVISO_ABERTURA_HORA[0], minute=AVISO_ABERTURA_HORA[1], second=0, microsecond=0)
        if not (hora <= agora < hora + timedelta(hours=6)):
            return
        if self.ligado is None or self.ligado > hora:
            return                       # robô não estava ligado antes das 00:15: não procurou
        dia_txt = agora.strftime("%d/%m/%Y")
        k = f"AVISO_SEM_ABERTURA|{dia_txt}"
        if k in self.estado.setdefault("enviadas", {}) or self.falhas.get(k, 0) >= MAX_TENTATIVAS:
            return
        if self.estado.get("aberturas", {}).get(dia_txt):
            return                       # abertura lida: tudo certo
        if self.driver is not None and time.time() - self.ultima_abertura > 60:
            self.ultima_abertura = time.time()       # última procura antes de avisar
            try:
                self.verificar_abertura()
            except Exception as e:
                if "invalid session" in str(e).lower() or "no such window" in str(e).lower():
                    raise
            if self.estado.get("aberturas", {}).get(dia_txt):
                return
        motivo = ("o email foi encontrado mas o robô não conseguiu ler a tabela de tanques"
                  if dia_txt in self._avisos_abertura else
                  "o email \"[Externo] ABERTURA DE TANQUES - " + dia_txt + "\" não foi encontrado no Outlook")
        self.log(f"AVISO: sem abertura de tanques de {dia_txt} ({motivo}) - avisando {EMAIL_AVISO_ABERTURA}")
        assunto = f"[Robô Laudo] Sem abertura de tanques - {dia_txt}"
        corpo = (f"{saudacao(agora)},\n\n"
                 f"O robô de pedido de laudo não recebeu a abertura de tanques do dia {dia_txt}: {motivo}.\n\n"
                 f"Por isso não foi feita a comparação dos tanques de S10, S500 e Gasolina com o dia anterior. "
                 f"Se algum tanque do pool mudou, peça o laudo manualmente ou use o botão \"Enviar agora\" no robô.\n\n"
                 "Atenciosamente,")
        reg = {"quando": agora.strftime("%d/%m/%Y %H:%M"), "tanque": "-", "produto": "AVISO SEM ABERTURA",
               "hora": "", "tq_base": "-"}
        try:
            enviar_email(self.driver, EMAIL_AVISO_ABERTURA, [], assunto, corpo, log=self.log)
        except Exception as e:
            self.falhas[k] = self.falhas.get(k, 0) + 1
            reg["status"] = f"FALHA: {e}"
            self.log(f"Falha ao enviar aviso de abertura ({self.falhas[k]}ª vez): {e}")
            self.fila.put(("enviado", reg))
            return
        reg["status"] = f"Aviso enviado -> {EMAIL_AVISO_ABERTURA}"
        self.estado["enviadas"][k] = datetime.now().isoformat()
        self.estado["historico"].append(reg)
        salvar_estado(self.estado)
        self.fila.put(("enviado", reg))
        self.fila.put(("alerta", f"Sem abertura de tanques {dia_txt[:5]} - responsável avisado"))

    def ler_abertura(self, data_txt, usar_cache=True):
        """Tanques do pool na abertura do dia 'dd/mm/aaaa' (cache no estado ou lendo o Outlook).
        usar_cache=False relê o email (pega ERRATA que chegue depois)."""
        cache = self.estado.setdefault("aberturas", {})
        if usar_cache and data_txt in cache:
            return cache[data_txt]
        itens = buscar_viradas(self.driver, TERMO_ABERTURA)
        if not itens:
            return None
        texto = ler_email_da_lista(self.driver, ["ABERTURA DE TANQUE", data_txt], preferir="ERRATA")
        if texto is None:
            return None
        tanques = parse_abertura(texto)
        if not tanques:
            if data_txt not in self._avisos_abertura:
                self._avisos_abertura.add(data_txt)
                self.log(f"Abertura {data_txt}: email achado mas não consegui ler a tabela "
                         f"(se a tabela vier como imagem, o robô não lê)")
            return cache.get(data_txt)
        cache[data_txt] = tanques
        salvar_estado(self.estado)
        return tanques

    def verificar_abertura(self):
        agora = datetime.now()
        # 23:25 em diante -> abertura de amanhã; resto do dia -> abertura de hoje
        dia = agora + timedelta(days=1) if (agora.hour, agora.minute) >= JANELA_ABERTURA[0] else agora
        hoje_txt, ontem_txt = dia.strftime("%d/%m/%Y"), (dia - timedelta(days=1)).strftime("%d/%m/%Y")
        hoje = self.ler_abertura(hoje_txt, usar_cache=False)
        if hoje is None:
            return                                    # ainda não chegou
        for p_, tq_ in hoje.items():
            self.registrar_pool(p_, tq_, datetime.strptime(hoje_txt, "%d/%m/%Y"))
        ontem = self.ler_abertura(ontem_txt)
        if ontem is None:
            self.log(f"Abertura {hoje_txt}: não achei a abertura de {ontem_txt} para comparar")
            return
        resumo = []
        for prod in ("GAS", "S10", "S500"):
            a, b = ontem.get(prod), hoje.get(prod)
            if not b:
                continue
            if a == b:
                resumo.append(f"{prod} TQ {b} (=)")
                continue
            resumo.append(f"{prod} TQ {a} -> {b} (MUDOU)")
            k = f"ABERTURA|{hoje_txt}|{prod}|{int(b)}"
            if k in self.estado.setdefault("enviadas", {}) or self.falhas.get(k, 0) >= MAX_TENTATIVAS:
                continue
            v = {"produto": PRODUTOS_ABERTURA[prod], "tanque": f"{int(b):02d}", "hora": ""}
            prod_rz, tanques_rz, origem_tq = escolher_tanques(v["produto"], self.passagem)
            if not tanques_rz:
                self.log(f"AGUARDANDO: abertura {hoje_txt} {prod} TQ {b} sem tanque Base ({prod_rz}) - "
                         f"tento de novo em {INTERVALO_ABERTURA_S // 60} min")
                self.fila.put(("alerta", f"Abertura {prod} TQ {b} aguardando tanque Base"))
                continue
            rz = [numero_tanque(n) for n in tanques_rz]
            ant = f"{int(a):02d}" if a else "?"
            linhas = [f"TQ {v['tanque']} (Pool) - {prod}"] + \
                     [f"TQ {t} (Base, em expedição) - {prod_rz}" for t in rz]
            assunto, corpo = montar_email_livre(
                f"Solicitação de laudo - TQ {v['tanque']} (Pool) + TQ {' + TQ '.join(rz)} (Base) - "
                f"{prod} - abertura {hoje_txt[:5]}",
                f"Na abertura de tanques do dia {hoje_txt}, o {prod} passou do TQ {ant} "
                f"para o TQ {v['tanque']} (Pool).", linhas)
            self.log(f"Abertura {hoje_txt}: {prod} mudou TQ {a} -> {b} | Base {'+'.join(rz)} ({origem_tq}) -> enviando")
            reg = {"quando": agora.strftime("%d/%m/%Y %H:%M"), "tanque": v["tanque"],
                   "produto": f"{prod} (abertura)", "hora": "", "tq_base": " + ".join(rz)}
            if self.enviar_pedido(assunto, corpo, reg, k):
                self.estado["enviadas"][k] = datetime.now().isoformat()
        chave_log = f"{hoje_txt}|{'|'.join(resumo)}"
        if resumo and getattr(self, "_ultimo_resumo_abertura", None) != chave_log:
            self._ultimo_resumo_abertura = chave_log
            self.log(f"Abertura {hoje_txt} x {ontem_txt}: " + "  |  ".join(resumo))
        salvar_estado(self.estado)

    def processar(self, v, k):
        self.atualizar_passagem()   # leitura fresca da planilha na hora da virada
        # tanque expedindo do TURNO ATUAL (o email de virada acabou de chegar)
        prod, tanques_rz, origem_tq = escolher_tanques(v["produto"], self.passagem)
        if not prod or not tanques_rz:
            # NUNCA envia só o produto puro: segura e tenta de novo até ter o tanque Base
            motivo = (f"produto '{v['produto']}' sem regra de composto" if not prod
                      else f"tanque de {prod} não achado na passagem nem na reserva da tela")
            if k.startswith("manual-"):
                self.log(f"Envio manual NÃO feito - {motivo}. Escolha o tanque na Reserva e clique de novo.")
                self.fila.put(("alerta", f"Envio manual TQ {v['tanque']} sem tanque Base"))
                return
            primeira = k not in self.pendentes
            self.pendentes[k] = time.time()
            if primeira:
                self.log(f"AGUARDANDO: virada TQ {v['tanque']} {v['produto']} {v['hora']} NÃO enviada - {motivo}. "
                         f"Tento de novo a cada {RETENTAR_PENDENTE_S // 60} min (ou escolha o tanque na reserva).")
                self.fila.put(("enviado", {
                    "quando": datetime.now().strftime("%d/%m/%Y %H:%M"), "tanque": v["tanque"],
                    "produto": v["produto"], "hora": v["hora"], "tq_base": "?",
                    "status": f"AGUARDANDO - {motivo}"}))
                self.fila.put(("alerta", f"Virada TQ {v['tanque']} {v['produto']} aguardando tanque Base"))
            return
        self.pendentes.pop(k, None)
        self.log(f"{prod}: TQ {' + '.join(numero_tanque(n) for n in tanques_rz)} (fonte: {origem_tq})")
        assunto, corpo = montar_email(v, tanques_rz)
        rz_txt = " + ".join(numero_tanque(n) for n in tanques_rz) or "-"
        self.log(f"Virada: TQ {v['tanque']} {v['produto']} {v['hora']} | Base expedindo: {rz_txt} -> enviando")
        reg = {
            "quando": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "tanque": v["tanque"], "produto": v["produto"], "hora": v["hora"],
            "tq_base": rz_txt,
        }
        if not self.enviar_pedido(assunto, corpo, reg, k):
            return  # não grava a chave -> tenta de novo
        self.estado["chaves"][k] = datetime.now().isoformat()
        self.estado.setdefault("enviadas", {})[self.chave_virada(v)] = datetime.now().isoformat()
        self.registrar_pool(v["produto"], v["tanque"], data_virada(v))
        salvar_estado(self.estado)
        self.log(f"Pedido de laudo enviado: TQ {v['tanque']} (Pool) + TQ {rz_txt} (Base).")
        if not self.pendentes:
            self.fila.put(("alerta_ok", None))


# ------------------------------ tela robô ---------------------------------
TELA_ROBO_AUTOMATICA = True    # PC dedicado: abre a tela cheia assim que começar a monitorar


class TelaRobo(tk.Toplevel):
    """Tela cheia de aviso 'PC ROBÔ - NÃO DESLIGAR'. O robô continua rodando por trás."""
    FUNDO = "#14081c"

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("PC ROBÔ - Monitor Virada de Tanque")
        self.configure(bg=self.FUNDO)
        self.geometry(f"{self.winfo_screenwidth()}x{self.winfo_screenheight()}+0+0")
        try:
            self.attributes("-fullscreen", True)
        except tk.TclError:
            pass
        self.attributes("-topmost", True)
        self.lift()
        self.protocol("WM_DELETE_WINDOW", self.sair)
        self.bind("<Escape>", lambda e: self.sair())

        f = tk.Frame(self, bg=self.FUNDO)
        f.place(relx=0.5, rely=0.5, anchor="center")

        tk.Label(f, text="PC ROBÔ EM OPERAÇÃO", bg=self.FUNDO, fg="#ffffff",
                 font=("Segoe UI", 44, "bold")).pack(pady=(0, 6))
        tk.Label(f, text="NÃO DESLIGAR  •  NÃO REINICIAR  •  NÃO FAZER LOGOFF  •  NÃO FECHAR O CHROME",
                 bg=self.FUNDO, fg="#ffffff", font=("Segoe UI", 20, "bold"), padx=24, pady=10).pack(pady=(0, 18))
        tk.Label(f, text="Monitor de Virada de Tanque  →  pedido automático de laudo ao laboratório",
                 bg=self.FUNDO, fg="#b9a6c9", font=("Segoe UI", 16)).pack(pady=(0, 28))

        self.lbl_status = tk.Label(f, bg=self.FUNDO, fg="#60c060", font=("Segoe UI", 30, "bold"))
        self.lbl_status.pack()
        self.lbl_relogio = tk.Label(f, bg=self.FUNDO, fg="#ffffff", font=("Consolas", 26))
        self.lbl_relogio.pack(pady=(4, 20))
        self.lbl_modo = tk.Label(f, bg=self.FUNDO, font=("Segoe UI", 16, "bold"))
        self.lbl_modo.pack(pady=(0, 14))
        self.lbl_passagem = tk.Label(f, bg=self.FUNDO, fg="#e0d4ea", font=("Segoe UI", 16),
                                     wraplength=int(self.winfo_screenwidth() * 0.85), justify="center")
        self.lbl_passagem.pack(pady=(0, 8))
        self.lbl_ultimo = tk.Label(f, bg=self.FUNDO, fg="#e0d4ea", font=("Segoe UI", 16),
                                   wraplength=int(self.winfo_screenwidth() * 0.85), justify="center")
        self.lbl_ultimo.pack(pady=(0, 8))
        self.lbl_alerta = tk.Label(f, bg=self.FUNDO, fg="#ff6b5b", font=("Segoe UI", 18, "bold"),
                                   wraplength=1200)
        self.lbl_alerta.pack(pady=(14, 0))

        rod = tk.Frame(self, bg=self.FUNDO)
        rod.pack(side="bottom", fill="x", pady=12)
        tk.Label(rod, text="Responsável: Operações  •  ESC ou botão ao lado para ver o log",
                 bg=self.FUNDO, fg="#6f5a80", font=("Segoe UI", 11)).pack(side="left", padx=20)
        tk.Button(rod, text="Ver detalhes / log", command=self.sair, bg="#2a1838", fg="#ffffff",
                  relief="flat", font=("Segoe UI", 11), padx=12).pack(side="right", padx=20)
        self.atualizar()

    def atualizar(self):
        if not self.winfo_exists():
            return
        a = self.app
        status = a.lbl_status.cget("text")
        ok = status.startswith("Monitorando")
        self.lbl_status.config(text=("●  " if ok else "▲  ") + status,
                               fg="#60c060" if ok else "#f0b030")
        self.lbl_relogio.config(text=f"{datetime.now():%d/%m/%Y  %H:%M:%S}   |   {a.lbl_tick.cget('text')}")
        teste = a.var_teste.get()
        self.lbl_modo.config(text=(f"MODO TESTE - envia para {EMAIL_TESTE}" if teste
                                   else f"PRODUÇÃO - envia para {EMAIL_LABORATORIO}"),
                             fg="#f0b030" if teste else "#60c060")
        self.lbl_passagem.config(text="Passagem de turno:  " + a.lbl_passagem.cget("text"))
        itens = a.tree.get_children()
        if itens:
            q, tqp, tqr, prod, hora, st = a.tree.item(itens[0], "values")
            hoje = datetime.now().strftime("%d/%m/%Y")
            n = sum(1 for i in itens if str(a.tree.item(i, "values")[0]).startswith(hoje)
                    and str(a.tree.item(i, "values")[5]).startswith("Enviado"))
            self.lbl_ultimo.config(text=f"Último pedido: {q}  —  TQ {tqp} {prod} {hora} + TQ {tqr} Base  →  "
                                        f"{st.split(' (')[0]}\nPedidos enviados hoje: {n}")
        else:
            self.lbl_ultimo.config(text="Nenhum pedido enviado ainda")
        self.lbl_alerta.config(text=a.lbl_alerta.cget("text"))
        self.after(1000, self.atualizar)

    def sair(self):
        self.app.tela = None
        self.app.deiconify()
        self.destroy()


# ------------------------------ interface ---------------------------------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Monitor Virada de Tanque -> Laboratório")
        self.geometry("900x620")
        self.fila = queue.Queue()
        self.monitor = None

        topo = ttk.Frame(self, padding=8)
        topo.pack(fill="x")
        self.lbl_status = ttk.Label(topo, text="Parado", font=("Segoe UI", 12, "bold"))
        self.lbl_status.pack(side="left")
        self.lbl_tick = ttk.Label(topo, text="")
        self.lbl_tick.pack(side="left", padx=12)
        self.var_teste = tk.BooleanVar(value=MODO_TESTE)
        self.chk_teste = tk.Checkbutton(topo, text=f"Modo teste (envia para {EMAIL_TESTE})",
                                        variable=self.var_teste, command=self.trocar_modo,
                                        font=("Segoe UI", 9, "bold"))
        self.chk_teste.pack(side="left", padx=12)
        self.after(0, self.pintar_modo)
        ttk.Button(topo, text="Diagnóstico", command=self.diagnostico).pack(side="right", padx=4)
        ttk.Button(topo, text="Tela Robô", command=self.abrir_tela_robo).pack(side="right", padx=4)
        self.tela = None
        self._tela_auto_feita = False
        ttk.Button(topo, text="Parar", command=self.parar).pack(side="right")
        ttk.Button(topo, text="Iniciar", command=self.iniciar).pack(side="right", padx=4)

        self.lbl_alerta = tk.Label(self, text="", fg="#c0392b", font=("Segoe UI", 10, "bold"))
        self.lbl_alerta.pack(fill="x", padx=8)

        ps = ttk.LabelFrame(self, text="Tanque Base EXPEDINDO - lido da Passagem de Turno (SharePoint)",
                            padding=8)
        ps.pack(fill="x", padx=8)
        self.lbl_passagem = ttk.Label(ps, text="aguardando leitura da planilha...",
                                      font=("Segoe UI", 10, "bold"))
        self.lbl_passagem.pack(side="left")

        exp = ttk.LabelFrame(self, text="Reserva - usada só se a planilha falhar ou estiver vazia", padding=8)
        exp.pack(fill="x", padx=8, pady=(6, 0))
        salvo = carregar_expedindo()
        self.combos, self.lbls_exp = {}, {}
        for prod, rotulo in (("BIODIESEL B100", "B100 (vai com S10 / S500):"),
                             ("ETANOL ANIDRO", "Anidro (vai com Gasolina):")):
            ttk.Label(exp, text=rotulo, font=("Segoe UI", 10, "bold")).pack(side="left")
            cb = ttk.Combobox(exp, values=[""] + TANQUES_BASE[prod], width=6, state="readonly",
                              font=("Segoe UI", 11, "bold"))
            cb.set((salvo.get(prod) or {}).get("tanque", ""))
            cb.bind("<<ComboboxSelected>>", lambda e, p=prod: self.trocar_expedindo(p))
            cb.pack(side="left", padx=(4, 2))
            lbl = ttk.Label(exp, text="", foreground="#666")
            lbl.pack(side="left", padx=(0, 18))
            self.combos[prod], self.lbls_exp[prod] = cb, lbl
        self.pintar_expedindo()

        teste = ttk.LabelFrame(self, text="Testar leitura de assunto (não envia nada)", padding=8)
        teste.pack(fill="x", padx=8, pady=(8, 0))
        self.ent_teste = ttk.Entry(teste)
        self.ent_teste.insert(0, "RE: [Externo] Virada de OD S10 p/ TQ 09 às 04h30 - TODOS")
        self.ent_teste.pack(side="left", fill="x", expand=True)
        ttk.Button(teste, text="Testar", command=self.testar).pack(side="left", padx=4)
        ttk.Button(teste, text="Enviar agora", command=self.enviar_manual).pack(side="left")

        cols = ("quando", "tanque", "tq_base", "produto", "hora", "status")
        self.tree = ttk.Treeview(self, columns=cols, show="headings", height=8)
        for c, t, w in zip(cols, ("Enviado em", "TQ Pool", "TQ Base", "Produto", "Hora virada", "Status"),
                           (130, 70, 90, 110, 90, 330)):
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w, anchor="w")
        self.tree.pack(fill="x", padx=8, pady=8)

        self.txt = tk.Text(self, height=14, font=("Consolas", 9))
        self.txt.pack(fill="both", expand=True, padx=8, pady=(0, 8))

        for r in (carregar_estado() or {}).get("historico", [])[-50:]:
            self.inserir_hist(r)

        self.protocol("WM_DELETE_WINDOW", self.sair)
        self.after(300, self.loop_fila)
        self.after(1000, self.iniciar)  # PC dedicado: já começa monitorando

    def escrever_log(self, msg):
        linha = f"[{datetime.now():%d/%m %H:%M:%S}] {msg}"
        self.txt.insert("end", linha + "\n")
        self.txt.see("end")
        try:
            with open(ARQ_LOG, "a", encoding="utf-8") as f:
                f.write(linha + "\n")
        except Exception:
            pass

    def inserir_hist(self, r):
        self.tree.insert("", 0, values=(r["quando"], r["tanque"], r.get("tq_base", "-"),
                                        r["produto"], r["hora"], r["status"]))

    def loop_fila(self):
        try:
            while True:
                tipo, dado = self.fila.get_nowait()
                if tipo == "log":
                    self.escrever_log(dado)
                elif tipo == "status":
                    self.lbl_status.config(text=dado)
                    if dado == "Monitorando" and TELA_ROBO_AUTOMATICA and not self._tela_auto_feita:
                        self._tela_auto_feita = True
                        self.abrir_tela_robo()
                elif tipo == "tick":
                    self.lbl_tick.config(text=f"última leitura {dado}")
                elif tipo == "enviado":
                    self.inserir_hist(dado)
                elif tipo == "alerta_ok":
                    self.lbl_alerta.config(text="")
                elif tipo == "alerta":
                    self.bell()
                    self.lbl_alerta.config(text=f"⚠ {dado} - escolha o tanque na Reserva se a planilha não tiver")
                elif tipo == "passagem":
                    hora, p = dado
                    partes = []
                    for prod, nome in (("BIODIESEL B100", "B100"), ("ETANOL ANIDRO", "Anidro"),
                                       ("ETANOL HIDRATADO", "Hidratado")):
                        if prod == "ETANOL HIDRATADO":
                            tq_eh = tanques_passagem(prod, p)
                            partes.append(f"Hidratado: TQ {' + '.join(tq_eh)}" if tq_eh else "Hidratado: -")
                            continue
                        _p, tqs, origem_tq = escolher_tanques("OD S10" if nome == "B100" else "GAS", p)
                        info = p.get(prod)
                        if tqs and origem_tq.startswith("passagem"):
                            partes.append(f"{nome}: TQ {' + '.join(numero_tanque(t) for t in tqs)} "
                                          f"(aba {info['aba']}, {info['turno']}º turno)")
                        elif tqs:
                            partes.append(f"{nome}: TQ {' + '.join(numero_tanque(t) for t in tqs)} (reserva)")
                        else:
                            partes.append(f"{nome}: NÃO ACHADO - escolha na reserva")
                    self.lbl_passagem.config(text=f"[{hora}]   " + "    |    ".join(partes))
        except queue.Empty:
            pass
        self.after(300, self.loop_fila)

    def iniciar(self):
        if self.monitor and self.monitor.is_alive():
            return
        self.monitor = Monitor(self.fila, modo_teste=self.var_teste.get())
        self.monitor.start()

    def parar(self):
        if self.monitor:
            self.monitor.parar.set()

    def diagnostico(self):
        if self.monitor and self.monitor.is_alive():
            self.monitor.pedir_diag.set()
            self.escrever_log("Diagnóstico pedido - sai no log em até 20 s.")
        else:
            self.escrever_log("Monitor não está rodando - clique em Iniciar.")

    def enviar_manual(self):
        v = extrair_virada(self.ent_teste.get())
        if not v:
            messagebox.showwarning("Enviar agora", "Não reconheci uma virada nesse texto.")
            return
        if not (self.monitor and self.monitor.is_alive() and self.monitor.estado is not None):
            messagebox.showwarning("Enviar agora", "O monitor ainda não está em 'Monitorando'.")
            return
        para = EMAIL_TESTE if self.var_teste.get() else EMAIL_LABORATORIO
        if messagebox.askyesno("Enviar agora",
                               f"Enviar pedido de laudo da virada\nTQ {v['tanque']} {v['produto']} {v['hora']}\n"
                               f"para {para}?\n\n(vai junto o tanque Base expedindo da Passagem de Turno)"):
            self.monitor.pedido_manual = v
            self.escrever_log("Envio manual na fila - sai em até 20 s.")

    def testar(self):
        v = extrair_virada(self.ent_teste.get())
        if not v:
            messagebox.showwarning("Teste", "Não reconheci uma virada nesse texto.")
            return
        pas = self.monitor.passagem if self.monitor else {}
        prod, tanques_rz, origem_tq = escolher_tanques(v["produto"], pas)
        assunto, corpo = montar_email(v, tanques_rz)
        if prod and not tanques_rz:
            aviso = f"\n\n({prod} não achado na passagem nem na reserva - vai só o pool)"
        else:
            aviso = f"\n\n(tanque Base: {origem_tq})" if origem_tq else ""
        para = EMAIL_TESTE if self.var_teste.get() else EMAIL_LABORATORIO
        messagebox.showinfo("Teste", f"Para: {para}\nAssunto: {assunto}\n\n{corpo}{aviso}")

    def trocar_expedindo(self, prod):
        d = carregar_expedindo()
        tq = self.combos[prod].get()
        if tq:
            d[prod] = {"tanque": tq, "quando": datetime.now().strftime("%d/%m %H:%M")}
            self.escrever_log(f"{prod} expedindo agora: TQ {tq}")
        else:
            d.pop(prod, None)
            self.escrever_log(f"{prod}: nenhum tanque expedindo informado")
        salvar_expedindo(d)
        self.pintar_expedindo()

    def pintar_expedindo(self):
        d = carregar_expedindo()
        for prod, lbl in self.lbls_exp.items():
            info = d.get(prod)
            lbl.config(text=f"desde {info['quando']}" if info else "não informado",
                       foreground="#666" if info else "#c0392b")

    def pintar_modo(self):
        teste = self.var_teste.get()
        self.chk_teste.config(fg="#c0392b" if teste else "#2e7d32",
                              text=(f"MODO TESTE (envia para {EMAIL_TESTE})" if teste
                                    else f"PRODUÇÃO (envia para {EMAIL_LABORATORIO})"))

    def trocar_modo(self):
        teste = self.var_teste.get()
        if not teste and not messagebox.askyesno(
                "Produção", f"Passar a enviar de verdade para {EMAIL_LABORATORIO}?"):
            self.var_teste.set(True)
            teste = True
        if self.monitor:
            self.monitor.modo_teste = teste
        self.pintar_modo()
        self.escrever_log("Modo TESTE ligado" if teste else "Modo PRODUÇÃO ligado - enviando para o laboratório")

    def abrir_tela_robo(self):
        if self.tela is None:
            self.tela = TelaRobo(self)
            self.withdraw()   # janela de detalhes some; volta com ESC / "Ver detalhes"

    def sair(self):
        if not messagebox.askyesno("Fechar o robô",
                                   "Fechar o Monitor de Virada?\n\nOs pedidos de laudo ao laboratório vão PARAR "
                                   "até alguém abrir o robô de novo."):
            return
        self.parar()
        if self.monitor:
            self.monitor.fechar_driver()   # fecha o Chrome para não travar o perfil na próxima vez
        self.after(500, self.destroy)


def manter_pc_acordado(ligar=True):
    """Impede o Windows de suspender/desligar a tela por inatividade enquanto o robô roda
    (não precisa de administrador). Bloqueio forçado pela política da empresa pode continuar."""
    try:
        import ctypes
        ES_CONTINUOUS, ES_SYSTEM_REQUIRED, ES_DISPLAY_REQUIRED = 0x80000000, 0x1, 0x2
        flags = ES_CONTINUOUS | (ES_SYSTEM_REQUIRED | ES_DISPLAY_REQUIRED if ligar else 0)
        ctypes.windll.kernel32.SetThreadExecutionState(flags)
    except Exception:
        pass


def instancia_unica(porta=47321):
    """Impede rodar dois monitores ao mesmo tempo (duplicaria pedidos e rascunhos)."""
    s_ = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        s_.bind(("127.0.0.1", porta))
        s_.listen(1)
        return s_
    except OSError:
        return None


if __name__ == "__main__":
    _trava = instancia_unica()
    if _trava is None:
        r = tk.Tk(); r.withdraw()
        messagebox.showerror("Monitor Virada", "O monitor já está aberto em outra janela.\n"
                             "Feche a outra (e o Chrome dela) antes de abrir de novo.")
    else:
        manter_pc_acordado(True)
        try:
            App().mainloop()
        finally:
            manter_pc_acordado(False)
